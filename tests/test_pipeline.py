"""Pipeline v1.1: extract → detect → apply (con revisión) → verify. Ejecutar: .venv/bin/pytest -q"""
import sys
from pathlib import Path
from types import SimpleNamespace

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from backend import engine, pipeline  # noqa: E402

ALL = [c for c in engine.ENTITY_TYPES if c != "FIRMA"]


def test_dos_apellidos_envueltos_con_coma_y_nombre_largo_tras_don():
    words=[['Marta',850,100,980,125],['Henríquez',100,158,235,184],
           ['Soto,',250,158,325,184],['chilena,',340,158,430,184]]
    page={'n':1,'width':1000,'height':2000,'words':words}
    finding=pipeline._finding('NOMBRE','Marta',page=1,boxes=[words[0][1:]])
    pipeline._extend_wrapped_names(page,[finding])
    assert finding['text']=='Marta\nHenríquez Soto'
    assert words[3][1:] not in finding['boxes']
    line=[['Don',100,100,145,125],['Florian',160,100,260,125],['Mauricio',280,100,390,125],
          ['Schneider',410,100,550,125],['Pérez,',570,100,670,125],['para',690,100,750,125]]
    candidates=pipeline._title_context_names({**page,'words':line})
    assert candidates[0]['text']=='Florian Mauricio Schneider Pérez'


def test_direccion_envuelta_incluye_oficina_y_marcador_con_punto():
    words=[['calle',750,100,810,125],['Los',825,100,865,125],['Aromos',880,100,985,125],
           ['N.°',100,170,150,195],['244,',170,170,225,195],['Oficina',240,170,320,195],['303,',330,170,385,195],['comuna',400,170,500,195]]
    findings=pipeline._wrapped_street_addresses({'n':1,'words':words})
    assert findings[0]['boxes']==[w[1:] for w in words[:7]]


def test_segunda_lectura_limitada_a_franca_omitida(monkeypatch):
    from PIL import Image
    words=[['Inicio',10,600,100,630],['párrafo',110,600,210,630],
           ['Siguiente',10,750,110,780],['párrafo',120,750,230,780]]
    calls=[]
    def sparse(img,scale):
        calls.append((img.size,scale));return [('Recuperado',20,45,150,70)]
    monkeypatch.setattr(pipeline,'_sparse_ocr_words',sparse)
    img=Image.new('RGB',(1000,2000),'white')
    new,count=pipeline._supplement_scan_gaps(img,words,'balanced')
    assert count==1 and len(calls)==1
    assert calls[0][0][0]==1000 and calls[0][0][1]<200 and calls[0][1]==1
    assert new==[['Recuperado',20,664,150,689]]
    assert not pipeline._scan_gap_regions(words[:2],2000)


def test_direccion_con_numero_explicitamente_envuelto():
    words=[['calle',700,100,770,130],['Los',780,100,820,130],['Aromos',830,100,945,130],
           ['N°',100,142,130,172],['452,',140,142,195,172],['comuna',205,142,300,172]]
    page={'n':1,'width':1000,'height':2000,'words':words}
    findings=pipeline._detect_page_text_findings(page,words,['DIRECCION'])
    assert any(f['boxes']==[w[1:] for w in words[:5]] for f in findings)
    words[3][0]='año'
    assert not pipeline._wrapped_street_addresses(page)


def test_cajas_pdf_cubren_apellido_final_de_lista_de_personas():
    tokens = "Rene Valencia Uribe y Diego Vergara Arriagada, con fecha".split()
    words = [[token, 10 + i * 100, 100, 100 + i * 100, 125]
             for i, token in enumerate(tokens)]
    page = {"n": 2, "width": 1500, "height": 2000, "words": words}
    findings = pipeline._detect_page_text_findings(page, words, ["NOMBRE"])
    assert [f["text"] for f in findings] == ["Rene Valencia Uribe", "Diego Vergara Arriagada"]
    assert findings[1]["boxes"] == [w[1:] for w in words[4:7]]
    assert not any(words[3][1:] in f["boxes"] for f in findings)  # conservar «y»


def test_apellido_en_linea_siguiente_no_arrastra_prosa():
    for ending, continuation, expected in [
        ("Vergara", "Arriagada", "Diego Vergara\nArriagada"),
        ("Vergara", "arriagada", "Diego Vergara\narriagada"),
        ("Arriagada", "antecedentes", "Diego Arriagada"),
        ("Arriagada,", "Nueva", "Diego Arriagada"),
    ]:
        words = [["Diego", 100, 100, 170, 130], [ending, 180, 100, 290, 130],
                 [continuation, 140, 135, 280, 165]]
        page = {"n": 2, "width": 1000, "height": 2000, "words": words}
        finding = pipeline._finding("NOMBRE", "Diego " + ending.rstrip(","),
                                    page=2, boxes=[w[1:] for w in words[:2]])
        pipeline._extend_wrapped_names(page, [finding])
        assert finding["text"] == expected
        assert (words[2][1:] in finding["boxes"]) == ("\n" in expected)


def test_nombre_destacado_en_cv_incluye_apellido_alineado():
    words = [['Marta', 100, 80, 260, 135], ['Henríquez', 100, 150, 330, 205],
             ['Profesional', 100, 230, 250, 250], ['experiencia', 270, 230, 420, 250],
             ['Formación', 100, 260, 240, 280], ['académica', 270, 260, 400, 280]]
    page = {'n': 1, 'width': 1000, 'height': 2000, 'words': words}
    finding = pipeline._finding('NOMBRE', 'Marta', page=1, boxes=[words[0][1:]])
    pipeline._extend_wrapped_names(page, [finding])
    assert finding['text'] == 'Marta\nHenríquez'
    assert finding['boxes'] == [words[0][1:], words[1][1:]]


def test_nombre_no_se_extiende_al_cargo_de_registros():
    words = [['Ana', 100, 100, 145, 123], ['Pérez', 155, 100, 210, 123],
             ['DE', 125, 128, 155, 151], ['REGISTROS', 160, 128, 275, 151]]
    page = {'n': 1, 'width': 1000, 'height': 2000, 'words': words}
    finding = pipeline._finding('NOMBRE', 'Ana Pérez', page=1, boxes=[w[1:] for w in words[:2]])
    pipeline._extend_wrapped_names(page, [finding])
    assert finding['text'] == 'Ana Pérez'


def test_apellido_partido_y_nombre_poco_comun_antes_de_jueza():
    words = [['Resolvió,', 10, 100, 110, 130], ['Zylka', 120, 100, 190, 130],
             ['M', 200, 100, 225, 130], ['ó', 100, 100, 108, 130],
             ['üller', 225, 100, 290, 130], ['Ejemplo,', 300, 100, 385, 130],
             ['Jueza', 395, 100, 450, 130], ['Titular', 460, 100, 540, 130]]
    page = {'n': 1, 'width': 1000, 'height': 2000, 'words': words}
    findings = pipeline._detect_page_text_findings(page, words, ['NOMBRE'])
    assert any(f['text'] == 'Zylka Müller Ejemplo' for f in findings)
    assert all(words[3][1:] not in f['boxes'] for f in findings)


def test_apellido_al_principio_de_linea_en_lista_de_abogados():
    words = [['Benjamín', 700, 100, 810, 130], ['Ficticio', 825, 100, 940, 130],
             ['Prueba', 100, 140, 180, 170], ['y', 190, 140, 200, 170],
             ['Sergio', 210, 140, 280, 170], ['Soto', 290, 140, 350, 170]]
    page = {'n': 1, 'width': 1000, 'height': 2000, 'words': words}
    finding = pipeline._finding('NOMBRE', 'Benjamín Ficticio', page=1, boxes=[w[1:] for w in words[:2]])
    pipeline._extend_wrapped_names(page, [finding])
    assert finding['text'] == 'Benjamín Ficticio\nPrueba'
    assert finding['boxes'] == [w[1:] for w in words[:3]]


def test_caja_del_titulo_solapada_no_omite_primer_nombre():
    words = [['SR.', 10, 100, 50, 130], ['JOSE', 45, 100, 115, 130],
             ['MIGUEL', 120, 100, 220, 130], ['PRUEBA', 225, 100, 340, 130],
             ['EJEMPLO', 345, 100, 450, 130]]
    page = {'n': 1, 'width': 1000, 'height': 2000, 'words': words}
    candidates = pipeline._title_context_names(page)
    assert any(f['text'] == 'JOSE MIGUEL PRUEBA EJEMPLO' for f in candidates)
    assert all(words[0][1:] not in f['boxes'] for f in candidates)


def test_perfiles_de_analisis_controlan_las_pasadas_ocr_hibridas():
    page = SimpleNamespace(
        width=600,
        height=800,
        images=[{"x0": 0, "x1": 600, "top": 0, "bottom": 800}],
        lines=[],
        rects=[],
    )
    words = [{"text": "NOMBRE", "top": 10, "bottom": 20}]

    assert pipeline._hybrid_ocr_passes("fast", page, words) == (False, False)
    assert pipeline._hybrid_ocr_passes("balanced", page, words) == (True, False)
    assert pipeline._hybrid_ocr_passes("exhaustive", page, words) == (True, True)

    page.images = []
    assert pipeline._hybrid_ocr_passes("exhaustive", page, words) == (False, False)


def test_nivel_de_analisis_invalido_se_rechaza(tmp_path):
    src = tmp_path / "a.txt"
    src.write_text("texto", encoding="utf-8")
    try:
        pipeline.extract(src, tmp_path / "wd", "desconocido")
        raise AssertionError("debió rechazar el nivel desconocido")
    except engine.EngineError as exc:
        assert exc.code == "INVALID_ANALYSIS_LEVEL"


def test_txt_revision_rechazo_y_manual(tmp_path):
    src = tmp_path / "a.txt"
    src.write_text("RUT 12.345.678-5, correo a@b.cl. Proyecto Cóndor.", encoding="utf-8")
    ext = pipeline.extract(src, tmp_path / "wd")
    fs = pipeline.detect_findings(ext, ALL)
    assert {f["entity_code"] for f in fs} == {"RUT", "EMAIL"}
    email = next(f for f in fs if f["entity_code"] == "EMAIL")
    email["status"] = "rejected"  # el revisor lo rechaza
    text = ext["texts"]["txt"]
    i = text.index("Cóndor")
    fs.append({"id": "m1", "entity_code": "NOMBRE", "text": "Cóndor", "source": "manual", "status": "accepted",
               "page": None, "boxes": None, "location": "txt", "start": i, "end": i + 6, "score": 1.0})
    out = tmp_path / "a_out.txt"
    stats = pipeline.apply_findings(src, ext, fs, tmp_path / "wd", out)
    got = out.read_text()
    assert "12.345.678-5" not in got and "a@b.cl" in got and "Cóndor" not in got
    assert stats["RUT"] == 1 and stats["NOMBRE"] == 1 and "EMAIL" not in stats

    # verificador: ok con lo aplicado; detecta filtración si exigimos también lo rechazado
    assert pipeline.verify(out, ["12.345.678-5", "Cóndor"])["ok"]
    v = pipeline.verify(out, ["a@b.cl"])
    assert not v["ok"] and "a@b.cl" in v["checks"][0]["detail"]


def test_docx_extract_apply_verify(tmp_path):
    import docx
    d = docx.Document()
    d.add_paragraph("Sr. Pedro Soto, RUT 12.345.678-5")
    d.add_paragraph("Sin datos sensibles aquí")
    d.core_properties.author = "Secreto"
    src = tmp_path / "b.docx"
    d.save(src)
    ext = pipeline.extract(src, tmp_path / "wd")
    fs = pipeline.detect_findings(ext, ALL)
    assert any(f["entity_code"] == "RUT" and f["location"] == "body/p/0" for f in fs)
    out = tmp_path / "b_out.docx"
    pipeline.apply_findings(src, ext, fs, tmp_path / "wd", out)
    txt = docx.Document(str(out)).paragraphs[0].text
    assert "12.345.678-5" not in txt and "Pedro" not in txt
    v = pipeline.verify(out, ["12.345.678-5", "Pedro Soto"])
    assert v["ok"], v


def test_paged_png_manual_box_y_verify(tmp_path):
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (900, 300), "white")
    dr = ImageDraw.Draw(img)
    font = ImageFont.load_default(size=40)
    dr.text((20, 40), "RUT 12.345.678-5", fill="black", font=font)
    dr.text((20, 140), "Texto normal", fill="black", font=font)
    src = tmp_path / "c.png"
    img.save(src)
    wd = tmp_path / "wd"
    ext = pipeline.extract(src, wd)
    fs = pipeline.detect_findings(ext, ALL, wd)
    assert any(f["entity_code"] == "RUT" and f["boxes"] for f in fs)
    fs.append({"id": "m1", "entity_code": "NOMBRE", "text": "(manual)", "source": "manual", "status": "accepted",
               "page": 1, "boxes": [[15, 130, 400, 200]], "location": None, "start": None, "end": None, "score": 1.0})
    out = tmp_path / "c_out.png"
    # Cached previews can be cleaned up or interrupted without losing the
    # original; protection must still reconstruct and cover the actual text.
    (wd / ext['pages'][0]['image']).write_bytes(b'PNG-incompleto')
    pipeline.apply_findings(src, ext, fs, wd, out)
    assert pipeline.page_image_ok(wd / ext['pages'][0]['image'])
    v = pipeline.verify(out, ["12.345.678-5"])
    assert v["ok"], v  # re-OCR del protegido no encuentra el RUT


def test_imagen_de_pagina_atomica_y_regenerable(tmp_path):
    from PIL import Image

    src = tmp_path / "pagina.png"
    Image.new("RGB", (320, 180), "white").save(src)
    workdir = tmp_path / "work"
    extraction = pipeline.extract(src, workdir)
    page = workdir / extraction["pages"][0]["image"]
    assert pipeline.page_image_ok(page)

    page.write_bytes(b"\x89PNG\r\n\x1a\narchivo-incompleto")
    assert not pipeline.page_image_ok(page)
    assert pipeline.render_page(src, workdir, 1) == page
    assert pipeline.page_image_ok(page)
    assert pipeline.render_page(src, workdir, 2) is None
    assert list(page.parent.glob(".*.tmp")) == []


def test_ocr_complementario_agrega_solo_palabras_ausentes(monkeypatch):
    from PIL import Image

    native = [["JOSE", 100, 500, 170, 530], ["EJEMPLO", 300, 500, 390, 530]]
    monkeypatch.setattr(pipeline, "_ocr_words", lambda _img: [])
    monkeypatch.setattr(pipeline, "_sparse_ocr_words", lambda _img: [
        ("JOSE", 100, 220, 170, 250),
        ("MIGUEL", 180, 220, 275, 250),
        ("EJEMPLO", 300, 220, 390, 250),
        ("PABLINA", 100, 320, 200, 350),
    ])
    img = Image.new("RGB", (600, 1000), "white")
    supplement = pipeline._supplement_scanned_words(img, native, start_ratio=0.28)
    assert supplement == [
        ("MIGUEL", 180, 500, 275, 530),
        ("PABLINA", 100, 600, 200, 630),
    ]


def test_paged_respeta_lineas_y_ancla_los_tres_campos_ocr_ruidosos():
    words = [
        ["Nombre", 10, 10, 80, 30], ["de", 90, 10, 110, 30], ["persona:", 120, 10, 200, 30],
        ["Valentina", 210, 10, 300, 30], ["Caso", 310, 10, 350, 30], ["Fuentes", 360, 10, 430, 30],
        ["RUT:", 10, 55, 50, 75], ["11.111.111-1", 60, 55, 180, 75],
        ["Correo:", 10, 100, 80, 120], ["v.caso@example.test", 90, 100, 270, 120],
        ["Cuenta", 10, 145, 70, 165], ["corriente", 80, 145, 160, 165],
        ["N”", 170, 145, 195, 165], ["1234567890", 205, 145, 305, 165],
        ["Pasaporte", 10, 190, 90, 210], ["N*", 100, 190, 125, 210], ["PTEST1234", 135, 190, 225, 210],
    ]
    extraction = {"kind": "paged", "format": "png", "pages": [
        {"n": 1, "image": "pages/p1.png", "width": 500, "height": 300, "words": words}
    ], "ocr_pages": [1]}
    fs = pipeline.detect_findings(
        extraction, ["NOMBRE", "RUT", "EMAIL", "CUENTA_BANCARIA", "PASAPORTE"]
    )
    by_code = {f["entity_code"]: f for f in fs}
    assert set(by_code) == {"NOMBRE", "RUT", "EMAIL", "CUENTA_BANCARIA", "PASAPORTE"}
    assert by_code["NOMBRE"]["text"] == "Valentina Caso Fuentes"
    assert by_code["NOMBRE"]["boxes"] == [words[i][1:] for i in (3, 4, 5)]
    assert by_code["EMAIL"]["boxes"][0][2] > words[9][3]  # holgura para TLD subdimensionado por OCR
    assert by_code["CUENTA_BANCARIA"]["boxes"] == [words[13][1:]]
    assert by_code["PASAPORTE"]["boxes"] == [words[16][1:]]


def test_propagacion_de_nombre_respeta_limites_de_palabra():
    extraction = {"kind": "paged", "format": "png", "pages": [{
        "n": 1, "image": "pages/p1.png", "width": 800, "height": 300,
        "words": [
            ["Nombre:", 10, 10, 80, 30], ["Dora", 90, 10, 130, 30], ["Eralon", 140, 10, 200, 30],
            ["asegurando", 10, 60, 100, 80], ["COMERCIALIZADORA", 110, 60, 260, 80],
            ["ERALON", 270, 60, 330, 80],
        ],
    }], "ocr_pages": [1]}
    fs = pipeline.detect_findings(extraction, ["NOMBRE"])
    assert [f["text"] for f in fs] == ["Dora Eralon"]


def test_propagacion_nombre_tolera_primer_nombre_abreviado_por_ocr():
    extraction = {"kind": "paged", "format": "pdf", "pages": [{
        "n": 1, "image": "pages/p1.png", "width": 800, "height": 300,
        "words": [],
        "verification_words": [
            ["MIG", 100, 100, 140, 125], ["PRUEBA", 155, 100, 250, 125],
            ["EJEMPLO", 265, 100, 350, 125],
        ],
    }], "ocr_pages": []}
    seed = [
        {"entity_code": "NOMBRE", "text": "MIGUEL PRUEBA EJEMPLO",
         "page": 2, "boxes": [[1, 1, 2, 2]]},
        {"entity_code": "NOMBRE", "text": "MIGUEL PRUEBA EJEMPLO",
         "page": 3, "boxes": [[1, 1, 2, 2]]},
    ]
    extra = pipeline._propagate_fuzzy_three_part_names(seed, extraction)
    assert [finding["text"] for finding in extra] == ["MIG PRUEBA EJEMPLO"]


def test_titulo_ocr_no_convierte_una_palabra_institucional_en_nombre():
    page = {"n": 21, "width": 1500, "height": 2200, "words": [
        ["Sra", 1200, 300, 1240, 330], ["monic", 1250, 300, 1320, 330],
    ]}
    assert pipeline._title_context_names(page) == []
    page["words"] = [
        ["SR", 900, 300, 930, 330], ["CT", 940, 300, 970, 330],
        ["RB", 980, 300, 1010, 330], ["02", 1020, 300, 1050, 330],
    ]
    assert pipeline._title_context_names(page) == []


def test_excel_columnas_apellido_y_apellidos_son_nombre(tmp_path):
    import openpyxl
    src = tmp_path / "apellidos.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nombre", "Apellido Paterno", "Apellidos", "Apellido Materno"])
    ws.append(["Ana", "Pérez", "González Soto", "Rojas"])
    wb.save(src)
    extraction = pipeline.extract(src, tmp_path / "work")
    findings = pipeline.detect_findings(extraction, ["NOMBRE"])
    by_location = {f["location"]: f["text"] for f in findings if f.get("location")}
    assert by_location["Sheet!A2"] == "Ana"
    assert by_location["Sheet!B2"] == "Pérez"
    assert by_location["Sheet!C2"] == "González Soto"
    assert by_location["Sheet!D2"] == "Rojas"


def test_excel_apellido_completo_prevalece_sobre_hallazgo_parcial(tmp_path):
    import openpyxl

    src = tmp_path / "personas.xlsx"
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Nombre", "Apellido", "Edad"])
    ws.append(["Francisca", "Sepúlveda", 50])
    ws.append(["Felipe", "Rojas +56 9 1234 5678", 57])
    ws.append(["Ana", "Li", 21])
    wb.save(src)

    workdir = tmp_path / "work"
    extraction = pipeline.extract(src, workdir)
    findings = pipeline.detect_findings(extraction, ["NOMBRE", "TELEFONO"])
    surname_findings = [f for f in findings if f.get("location", "").startswith("Sheet!B")]
    assert [(f["location"], f["entity_code"], f["text"]) for f in surname_findings] == [
        ("Sheet!B2", "NOMBRE", "Sepúlveda"),
        ("Sheet!B3", "NOMBRE", "Rojas +56 9 1234 5678"),
        ("Sheet!B4", "NOMBRE", "Li"),
    ]

    dst = tmp_path / "protegido.xlsx"
    pipeline.apply_findings(src, extraction, findings, workdir, dst)
    result = openpyxl.load_workbook(dst).active
    assert result["B2"].value == "█" * len("Sepúlveda")
    assert not any(c.isalnum() for c in result["B3"].value)
    assert "█" in result["B3"].value
    assert result["B4"].value == "██"
    assert result["C2"].value == 50


def test_firma_heuristica(tmp_path):
    from PIL import Image, ImageDraw, ImageFont
    font = ImageFont.load_default(size=34)

    def base():
        img = Image.new("RGB", (1200, 800), "white")
        ImageDraw.Draw(img).text((100, 600), "Firma:", fill="black", font=font)
        return img

    # solo la línea de subrayado → no es firma
    img1 = base()
    ImageDraw.Draw(img1).line((260, 640, 700, 640), fill="black", width=4)
    w1 = engine._ocr_words(img1)
    assert engine._signature_boxes(img1, w1) == []

    # garabato grueso sobre la línea → firma detectada
    img2 = base()
    d = ImageDraw.Draw(img2)
    for i in range(6):
        d.arc((300 + i * 60, 560, 380 + i * 60, 650), 0, 360, fill="black", width=6)
    w2 = engine._ocr_words(img2)
    boxes = engine._signature_boxes(img2, w2)
    assert len(boxes) == 1, f"esperaba una sola firma, obtuvo {boxes}"
    x0, y0, x1, y1 = boxes[0]
    assert x1 - x0 > 100 and y1 - y0 > 30


def test_zip_solo_en_app():
    # el motor no conoce ZIP: la expansión es de la Plataforma (documents.py)
    assert ".zip" not in engine.SUPPORTED


def _docx_con_hipervinculo(path, texto, destino):
    """Párrafo con un hipervínculo externo, como los que arma Word al escribir un correo."""
    import docx
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    d = docx.Document()
    par = d.add_paragraph("Canal de reporte: ")
    rid = par.part.relate_to(destino, RT.HYPERLINK, is_external=True)
    enlace = OxmlElement("w:hyperlink")
    enlace.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    nodo = OxmlElement("w:t")
    nodo.text = texto
    run.append(nodo)
    enlace.append(run)
    par._p.append(enlace)
    d.save(str(path))
    return path


def _rels_xml(path) -> str:
    import zipfile
    with zipfile.ZipFile(path) as z:
        return "\n".join(z.read(n).decode("utf-8") for n in z.namelist() if n.endswith(".rels"))


def test_docx_no_deja_el_correo_en_el_destino_del_hipervinculo(tmp_path):
    """El texto tachado se ve bien, pero el `mailto:` vive en los .rels: abrir el .docx
    como ZIP lo dejaba a la vista aunque el documento se viera limpio."""
    src = _docx_con_hipervinculo(tmp_path / "c.docx", "soporte@example.test",
                                 "mailto:soporte@example.test")
    assert "soporte@example.test" in _rels_xml(src)      # el original sí lo trae

    ext = pipeline.extract(src, tmp_path / "wd")
    fs = pipeline.detect_findings(ext, ALL)
    assert any(f["entity_code"] == "EMAIL" for f in fs)
    out = tmp_path / "c_out.docx"
    pipeline.apply_findings(src, ext, fs, tmp_path / "wd", out, codes=ALL)
    assert "soporte@example.test" not in _rels_xml(out)

    import docx
    assert "soporte@example.test" not in "\n".join(p.text for p in docx.Document(str(out)).paragraphs)


def test_docx_conserva_el_enlace_si_no_se_pidio_tachar_correos(tmp_path):
    src = _docx_con_hipervinculo(tmp_path / "d.docx", "soporte@example.test",
                                 "mailto:soporte@example.test")
    ext = pipeline.extract(src, tmp_path / "wd")
    fs = pipeline.detect_findings(ext, ["RUT"])
    out = tmp_path / "d_out.docx"
    pipeline.apply_findings(src, ext, fs, tmp_path / "wd", out, codes=["RUT"])
    assert "soporte@example.test" in _rels_xml(out)


def test_docx_elimina_mailto_para_correo_agregado_manual(tmp_path):
    src = _docx_con_hipervinculo(tmp_path / "manual.docx", "soporte@example.test",
                                 "mailto:soporte@example.test")
    ext = pipeline.extract(src, tmp_path / "wd")
    manual = {"entity_code": "EMAIL", "text": "soporte@example.test", "status": "accepted"}
    out = tmp_path / "manual_out.docx"
    pipeline.apply_findings(src, ext, [manual], tmp_path / "wd", out, codes=["RUT"])
    assert "soporte@example.test" not in _rels_xml(out)
