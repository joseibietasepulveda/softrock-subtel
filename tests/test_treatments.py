"""Matriz de aceptación: cinco tratamientos × once formatos soportados."""
from __future__ import annotations

import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from backend import engine, pipeline  # noqa: E402


FORMATS = ["pdf", "doc", "docx", "xls", "xlsx", "txt", "jpg", "jpeg", "png", "tif", "tiff"]
TREATMENTS = list(pipeline.TREATMENTS)
ORIGINAL = "12.345.678-5"


def _make_source(root: Path, fmt: str) -> Path:
    if fmt == "txt":
        path = root / "entrada.txt"
        path.write_text(f"RUT {ORIGINAL}", encoding="utf-8")
        return path
    if fmt in ("docx", "doc"):
        import docx
        modern = root / "entrada.docx"
        document = docx.Document()
        document.add_paragraph(f"RUT {ORIGINAL}")
        document.save(modern)
        return modern if fmt == "docx" else engine._convert(modern, "doc", root)
    if fmt in ("xlsx", "xls"):
        import openpyxl
        modern = root / "entrada.xlsx"
        workbook = openpyxl.Workbook()
        workbook.active["A1"] = f"RUT {ORIGINAL}"
        workbook.save(modern)
        return modern if fmt == "xlsx" else engine._convert(modern, "xls", root)

    from PIL import Image, ImageDraw, ImageFont
    image = Image.new("RGB", (1600, 400), "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("DejaVuSans.ttf", 90)
    except OSError:
        font = ImageFont.load_default(size=90)
    draw.text((100, 120), f"RUT {ORIGINAL}", fill="black", font=font)
    if fmt == "pdf":
        import img2pdf
        png = root / "entrada.png"
        image.save(png)
        path = root / "entrada.pdf"
        path.write_bytes(img2pdf.convert(str(png)))
        return path
    path = root / f"entrada.{fmt}"
    if fmt in ("jpg", "jpeg"):
        image.save(path, format="JPEG", quality=95)
    elif fmt in ("tif", "tiff"):
        image.save(path, format="TIFF", compression="tiff_deflate")
    else:
        image.save(path, format="PNG")
    return path


def _output_pixels(path: Path):
    from PIL import Image
    if path.suffix.lower() == ".pdf":
        from pdf2image import convert_from_path
        return convert_from_path(str(path), dpi=pipeline.DPI)[0].convert("L")
    return Image.open(path).convert("L")


@pytest.mark.parametrize("treatment", TREATMENTS)
@pytest.mark.parametrize("fmt", FORMATS)
def test_tratamiento_cumple_objetivo_en_cada_formato(fmt, treatment, tmp_path):
    if fmt in ("doc", "xls") and not engine._soffice():
        pytest.fail("LibreOffice es obligatorio para validar DOC/XLS")
    source = _make_source(tmp_path, fmt)
    workdir = tmp_path / "work"
    extraction = pipeline.extract(source, workdir)
    findings = pipeline.detect_findings(extraction, ["RUT"], workdir)
    assert len(findings) == 1, (fmt, treatment, findings)

    output = tmp_path / f"salida.{fmt}"
    pseudonyms = pipeline.build_pseudonym_map(findings)
    stats = pipeline.apply_findings(source, extraction, findings, workdir, output,
                                    treatment, pseudonyms)
    assert output.exists() and output.stat().st_size > 0
    assert stats["treatment"] == treatment and stats["RUT"] == 1
    verification = pipeline.verify(output, [ORIGINAL])
    assert verification["ok"], (fmt, treatment, verification)

    expected = pipeline.replacement_for(findings[0], treatment, pseudonyms)
    if extraction["kind"] == "text":
        checkdir = tmp_path / "check"
        protected = pipeline.extract(output, checkdir)
        visible = "\n".join(protected["texts"].values())
        assert expected in visible, (fmt, treatment, expected, visible)
    else:
        # En formatos paginados se valida el efecto visual dentro de la caja:
        # negro opaco para tachado; fondo limpio con sustituto visible para el resto.
        page = extraction["pages"][0]
        image = _output_pixels(output)
        boxes = findings[0]["boxes"]
        x0, y0 = min(b[0] for b in boxes), min(b[1] for b in boxes)
        x1, y1 = max(b[2] for b in boxes), max(b[3] for b in boxes)
        sx, sy = image.width / page["width"], image.height / page["height"]
        crop = image.crop((round(x0 * sx), round(y0 * sy), round(x1 * sx), round(y1 * sy)))
        pixels = list(crop.getdata())
        dark = sum(value < 55 for value in pixels) / len(pixels)
        white = sum(value > 235 for value in pixels) / len(pixels)
        if treatment == "redact":
            assert dark > 0.9, (fmt, dark)
        else:
            assert white > 0.65 and dark > 0.002, (fmt, treatment, white, dark)


def test_ejemplos_y_consistencia_de_seudonimos():
    findings = [
        {"entity_code": "NOMBRE", "text": "Ana Pérez", "status": "accepted"},
        {"entity_code": "NOMBRE", "text": "ana  perez", "status": "accepted"},
        {"entity_code": "RUT", "text": ORIGINAL, "status": "accepted"},
    ]
    mapping = pipeline.build_pseudonym_map(findings)
    assert pipeline.replacement_for(findings[0], "pseudonymize", mapping) == "Persona-01"
    assert pipeline.replacement_for(findings[1], "pseudonymize", mapping) == "Persona-01"
    assert pipeline.replacement_for(findings[2], "pseudonymize", mapping) == "RUT-01"
    assert pipeline.replacement_for(findings[2], "anonymize", mapping) == "[RUT]"
    assert pipeline.replacement_for(findings[2], "mask_partial", mapping) == "12.3**.***-*"
    assert pipeline.replacement_for(findings[2], "mask_full", mapping) == "********"
    assert pipeline.replacement_for(findings[2], "redact", mapping) == "████████████"
@pytest.mark.parametrize('value,expected',[('X','*'),('Li','L*'),('Ana','An*'),('Luis','Lui*')])
def test_enmascaramiento_parcial_nunca_deja_completo_un_valor_corto(value,expected):
    finding=pipeline._finding('NOMBRE',value)
    assert pipeline.replacement_for(finding,'mask_partial')==expected

