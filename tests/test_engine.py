"""Chequeos mínimos del motor. Ejecutar: .venv/bin/pytest -q"""
import sys, zlib
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from backend import engine  # noqa: E402

ALL = list(engine.ENTITY_TYPES)
TXT = ("Solicitud de don Juan Pérez González, RUT 12.345.678-5, domicilio Av. Providencia 1234, Depto 56, Providencia, "
       "correo jperez@example.test, teléfono +56 9 8765 4321, nacido el 04/05/1980, patente BCDF-12, IP 10.0.0.7, "
       "pasaporte N° AB123456, cuenta corriente N° 12345678901, folio 2026-001234, tarjeta 4111 1111 1111 1111, "
       "www.subtel.gob.cl. La Subsecretaría de Telecomunicaciones y el Ministerio de Hacienda resuelven.")


def _codes(text, codes=ALL):
    return {s.code: text[s.start:s.end] for s in engine.detect(text, codes)}


def test_detecta_todos_los_tipos():
    found = _codes(TXT)
    assert found["RUT"] == "12.345.678-5"
    assert found["EMAIL"] == "jperez@example.test"
    assert found["TELEFONO"] == "+56 9 8765 4321"
    assert found["FECHA_NACIMIENTO"] == "04/05/1980"
    assert found["PATENTE"] == "BCDF-12"
    assert found["IP"] == "10.0.0.7"
    assert found["PASAPORTE"] == "AB123456"
    assert found["CUENTA_BANCARIA"] == "12345678901"
    assert found["FOLIO"] == "2026-001234"
    assert found["TARJETA"] == "4111 1111 1111 1111"
    assert found["WEB_REDES"] == "www.subtel.gob.cl"
    assert found["NOMBRE"] == "Juan Pérez González"
    assert found["DIRECCION"] == "Av. Providencia 1234, Depto 56, Providencia"


def test_no_tacha_instituciones_ni_rut_invalido():
    text = "Subsecretaría de Telecomunicaciones, Ministerio de Hacienda, RUT 12.345.678-9 (dígito inválido)."
    assert _codes(text) == {}


def test_respeta_seleccion_de_entidades():
    found = _codes(TXT, ["RUT", "EMAIL"])
    assert set(found) == {"RUT", "EMAIL"}


def test_redact_text_conserva_longitud_y_espacios():
    new, spans = engine.redact_text("RUT 12.345.678-5 fin", ["RUT"])
    assert new == "RUT ████████████ fin" and len(spans) == 1


def test_ruido_ocr_en_correo_cuenta_y_pasaporte():
    correo = engine._OCR_AT_RE.sub("@", "valentina.caso.fuentes(Gexample.test")
    assert correo == "valentina.caso.fuentes@example.test"
    text = f"Correo: {correo}\nCuenta corriente N” 1234567890\nPasaporte N* PTEST1234"
    found = _codes(text, ["EMAIL", "CUENTA_BANCARIA", "PASAPORTE"])
    assert found == {
        "EMAIL": "valentina.caso.fuentes@example.test",
        "CUENTA_BANCARIA": "1234567890",
        "PASAPORTE": "PTEST1234",
    }
    unicode_mail = "sofía.ficticia.pérez@example.test"
    assert _codes(unicode_mail, ["EMAIL"]) == {"EMAIL": unicode_mail}


def test_reconstruccion_de_palabras_conserva_lineas_visuales():
    words = [
        ("Nombre", 10, 10, 80, 30), ("de", 90, 10, 110, 30), ("persona:", 120, 10, 200, 30),
        ("Valentina", 210, 10, 300, 30), ("Caso", 310, 10, 350, 30), ("Fuentes", 360, 10, 430, 30),
        ("RUT:", 10, 55, 50, 75), ("11.111.111-1", 60, 55, 180, 75),
    ]
    text, index = engine._words_text_map(words)
    assert "Fuentes\nRUT:" in text
    assert len(text) == len(index)
    names = [text[s.start:s.end] for s in engine.detect(text, ["NOMBRE"])]
    assert names == ["Valentina Caso Fuentes"]


def _pdf_with_text(path: Path, text: str):
    """PDF mínimo de una página con texto Helvetica (nativo, extraíble)."""
    stream = f"BT /F1 12 Tf 50 700 Td ({text}) Tj ET".encode("latin-1")
    objs = [b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Length %d >>stream\n" % len(stream) + stream + b"\nendstream"]
    out, offs = b"%PDF-1.4\n", []
    for i, o in enumerate(objs, 1):
        offs.append(len(out)); out += b"%d 0 obj\n" % i + o + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objs) + 1) + b"".join(b"%010d 00000 n \n" % o for o in offs)
    out += b"trailer\n<< /Size %d /Root 1 0 R /Info << /Author (Secreto) >> >>\nstartxref\n%d\n%%%%EOF\n" % (len(objs) + 1, xref)
    path.write_bytes(out)


def test_pdf_nativo_queda_sin_texto_ni_metadatos(tmp_path):
    import pdfplumber
    src, dst = tmp_path / "in.pdf", tmp_path / "out.pdf"
    _pdf_with_text(src, "Contacto: jperez@example.test RUT 12.345.678-5")
    stats = engine.redact_file(src, dst, ALL)
    assert stats["EMAIL"] == 1 and stats["RUT"] == 1 and stats["pages"] == 1
    with pdfplumber.open(dst) as pdf:
        assert not (pdf.pages[0].extract_text() or "").strip()   # rasterizado: nada extraíble
        assert not (pdf.metadata or {}).get("Author")


def test_docx_y_xlsx_y_txt(tmp_path):
    import docx, openpyxl
    d = docx.Document(); d.add_paragraph("Sr. Pedro Soto, RUT 12.345.678-5"); d.core_properties.author = "Secreto"
    d.save(tmp_path / "a.docx")
    engine.redact_file(tmp_path / "a.docx", tmp_path / "a_out.docx", ALL)
    out = docx.Document(str(tmp_path / "a_out.docx"))
    assert "12.345.678-5" not in out.paragraphs[0].text and "Pedro" not in out.paragraphs[0].text
    assert out.core_properties.author == ""

    wb = openpyxl.Workbook(); wb.active["A1"] = "correo: a@b.cl"; wb.save(tmp_path / "b.xlsx")
    engine.redact_file(tmp_path / "b.xlsx", tmp_path / "b_out.xlsx", ["EMAIL"])
    assert openpyxl.load_workbook(tmp_path / "b_out.xlsx").active["A1"].value == "correo: ██████"

    (tmp_path / "c.txt").write_text("Tel 9 8765 4321", encoding="utf-8")
    engine.redact_file(tmp_path / "c.txt", tmp_path / "c_out.txt", ["TELEFONO"])
    assert (tmp_path / "c_out.txt").read_text() == "Tel █ ████ ████"


def test_imagen_con_ocr(tmp_path):
    from PIL import Image, ImageDraw, ImageFont
    img = Image.new("RGB", (900, 200), "white")
    ImageDraw.Draw(img).text((20, 60), "RUT 12.345.678-5", fill="black", font=ImageFont.load_default(size=40))
    img.save(tmp_path / "i.png")
    stats = engine.redact_file(tmp_path / "i.png", tmp_path / "i_out.png", ["RUT"])
    assert stats["RUT"] == 1
    out = Image.open(tmp_path / "i_out.png").convert("L")
    assert sum(1 for p in out.getdata() if p == 0) > 2000   # hay un bloque negro


def test_ocr_letras_grandes_conserva_coordenadas_originales(monkeypatch):
    from PIL import Image
    import pytesseract

    calls = []
    def read(image, **kwargs):
        calls.append(image.size)
        if len(calls) == 1:
            return {"text": ["12.345.0/8-5"], "left": [20], "top": [100],
                    "width": [500], "height": [160]}
        return {"text": ["12.345.678-5"], "left": [6], "top": [30],
                "width": [150], "height": [48]}
    monkeypatch.setattr(pytesseract, "image_to_data", read)
    with Image.new("RGB", (1000, 500), "white") as image:
        words = engine._ocr_words(image)
    assert calls == [(1000, 500), (300, 150)]
    assert words == [("12.345.678-5", 20, 100, 520, 260)]


def test_ocr_texto_normal_no_repite_lectura(monkeypatch):
    from PIL import Image
    import pytesseract

    calls = []
    def read(image, **kwargs):
        calls.append(image.size)
        return {"text": ["12.345.678-5"], "left": [20], "top": [30],
                "width": [100], "height": [20]}
    monkeypatch.setattr(pytesseract, "image_to_data", read)
    with Image.new("RGB", (1000, 500), "white") as image:
        assert engine._ocr_words(image) == [("12.345.678-5", 20, 30, 120, 50)]
    assert len(calls) == 1


def test_correo_partido_por_el_troceo_en_palabras():
    """pdfplumber y Tesseract cortan la dirección al separar la página en palabras.
    Las variantes reproducen separaciones habituales usando direcciones ficticias."""
    variantes = [
        "soporte@example. test",     # pdfplumber separó el TLD
        "soporte @example.test",     # separó la arroba
        "soporte@\nexample.test",    # la celda cortó la línea después de la arroba
        "soporte@example\n.test",    # y aquí antes del punto
        "soporte arroba example.test",
        "soporte (at) example.test",
        "soporte\uff20example.test",  # arroba de ancho completo
    ]
    for correo in variantes:
        texto = f"Canal de reporte de incidentes: {correo} / +56 9 1234 5678."
        found = _codes(texto, ["EMAIL"])
        assert found.get("EMAIL") == correo, correo


def test_correo_bien_formado_no_arrastra_el_texto_vecino():
    """La pasada tolerante solo entra donde la estricta no encontró nada."""
    assert _codes("Punto 3 .\ncontacto@example.test fin", ["EMAIL"]) == {"EMAIL": "contacto@example.test"}
    assert _codes("Enviar a\ncontacto@example.test", ["EMAIL"]) == {"EMAIL": "contacto@example.test"}
    assert _codes("Nos juntamos at google.com el lunes", ["EMAIL"]) == {}
    assert _codes("Ver el anexo 2. Correo institucional del servicio.", ["EMAIL"]) == {}


def test_homoglifos_no_desplazan_los_offsets():
    texto = "RUT 12.345.678\u20115 y correo j\uff0eperez\uff20subtel\uff0egob.cl."
    found = _codes(texto, ["RUT", "EMAIL"])
    assert found["RUT"] == "12.345.678\u20115"
    assert found["EMAIL"] == "j\uff0eperez\uff20subtel\uff0egob.cl"


def test_ruido_ocr_acepta_cualquier_letra_entre_parentesis():
    for ruido in ("(W", "(a)", "(Q)", "(g", "[e]"):
        crudo = f"valentina.caso{ruido}example.test"
        assert engine._OCR_AT_RE.sub("@", crudo) == "valentina.caso@example.test", ruido


def test_url_partida_por_salto_de_linea():
    """Un CV de diseño parte la URL al ancho de la columna: la primera línea termina en
    `/` y el usuario queda en la siguiente. La continuación solo entra tras un carácter
    explícito de continuación; una URL que termina limpia no arrastra la línea siguiente."""
    t = "Portafolio: https://portfolio.example.test/\npersona_ficticia y más"
    assert _codes(t, ["WEB_REDES"]) == {"WEB_REDES": "https://portfolio.example.test/\npersona_ficticia"}
    assert _codes("Visite www.subtel.cl\nSantiago, marzo", ["WEB_REDES"]) == {"WEB_REDES": "www.subtel.cl"}
    assert _codes("Ver https://ofertas.example.test/oferta?id=\n8841 antes del cierre", ["WEB_REDES"]) == \
        {"WEB_REDES": "https://ofertas.example.test/oferta?id=\n8841"}


def test_rangos_de_anios_no_son_telefonos():
    """Las líneas de tiempo de un CV («2023\\n2013») calzaban con el patrón de teléfono:
    área 2 + 023 + 2013. Un teléfono real no llega partido por un salto de línea ni
    formado por dos años."""
    assert _codes("EXPERIENCIA\n2023\n2013 Diseñadora", ["TELEFONO"]) == {}
    assert _codes("Período 2022 2018 en la empresa", ["TELEFONO"]) == {}
    # los teléfonos legítimos siguen saliendo
    assert "TELEFONO" in _codes("Llamar al +56 9 1234 5678", ["TELEFONO"])
    assert "TELEFONO" in _codes("Fono (2) 2123 4567", ["TELEFONO"])
    assert "TELEFONO" in _codes("anexo 32 2345 6789", ["TELEFONO"])


def test_rut_partido_y_solo_persona_natural():
    assert _codes("Cédula 11.111.111-\n1", ["RUT"]) == {"RUT": "11.111.111-\n1"}
    assert _codes("RUT 12 345 678-5", ["RUT"]) == {"RUT": "12 345 678-5"}
    assert _codes("Proveedor RUT 76.543.210-3", ["RUT"]) == {}


def test_fecha_no_es_patente_y_patentes_con_punto():
    text = "FECHA 15 DE SEPTIEMBRE DE\n2025; patentes AB.1234 y BCDF.12; DE 2025."
    found = [(s.code, text[s.start:s.end]) for s in engine.detect(text, ["FECHA", "PATENTE"])]
    assert found == [
        ("FECHA", "15 DE SEPTIEMBRE DE\n2025"),
        ("PATENTE", "AB.1234"),
        ("PATENTE", "BCDF.12"),
    ]
    assert engine.detect("ruido 02.01.001 y 31.02.2004", ["FECHA"]) == []


def test_roles_judiciales_completos_no_son_telefonos():
    text = ("Causa Rol C-1234-2020; Causa Rol C-1235-2020; "
            "Causa Rol C-1236-2021; Causa Rol C-1237-2021")
    found = [(s.code, text[s.start:s.end]) for s in engine.detect(text, ["FOLIO", "TELEFONO"])]
    assert found == [
        ("FOLIO", "C-1234-2020"), ("FOLIO", "C-1235-2020"),
        ("FOLIO", "C-1236-2021"), ("FOLIO", "C-1237-2021"),
    ]


def test_nombres_en_mayusculas_dentro_de_encabezados():
    text = "DEPARTAMENTO MIGUEL PRUEBA EJEMPLO DE SALUD\nPABLINA EJEMPLO PRUEBA\nYANET PRUEBA EJEMPLO"
    names = [text[s.start:s.end] for s in engine.detect(text, ["NOMBRE"])]
    assert names == ["MIGUEL PRUEBA EJEMPLO", "PABLINA EJEMPLO PRUEBA", "YANET PRUEBA EJEMPLO"]


def _pagina(w=1600, h=900):
    from PIL import Image
    return Image.new("RGB", (w, h), "white")


def test_firma_sin_texto_de_anclaje():
    """Una plancha que es pura firma manuscrita (0 palabras de OCR) no disparaba la
    heurística, anclada a la palabra «Firma». Con página sin texto se busca la mancha
    de tinta en toda la página; score bajo, el humano decide."""
    from PIL import ImageDraw
    img = _pagina()
    d = ImageDraw.Draw(img)
    for i in range(0, 720, 30):  # trazos curvos delgados (densidad ~9%, la variante de trazo fino mide 4%)
        d.arc((200 + i // 3, 300, 700 + i, 620), start=i % 360, end=(i + 140) % 360, fill=(20, 40, 160), width=4)
    boxes = engine._signature_boxes(img, [])
    assert len(boxes) == 1, boxes
    x0, y0, x1, y1 = boxes[0]
    tinta = img.convert("L").point(lambda v: 255 if v < 100 else 0).getbbox()
    assert x0 <= tinta[0] and y0 <= tinta[1] and x1 >= tinta[2] - 1 and y1 >= tinta[3] - 1  # contiene los trazos


def test_firma_fallback_no_dispara_con_texto_ni_fotos():
    from PIL import ImageDraw
    # página con texto normal (≥5 palabras) y sin la palabra Firma: el fallback no entra
    con_texto = _pagina()
    palabras = [(t, 100 + i * 200, 100, 250 + i * 200, 130) for i, t in
                enumerate(["informe", "anual", "de", "gestión", "2025", "resumen"])]
    ImageDraw.Draw(con_texto).line((100, 400, 900, 700), fill="black", width=6)
    assert engine._signature_boxes(con_texto, palabras) == []
    # imagen densa tipo foto: la mancha es demasiado llena para ser un trazo
    foto = _pagina()
    ImageDraw.Draw(foto).rectangle((200, 200, 1300, 800), fill=(40, 40, 40))
    assert engine._signature_boxes(foto, []) == []


def test_nombres_ampliados_y_cargos_judiciales():
    cases = {
        "profesor Jean Pierre Prueba, autor": "Jean Pierre Prueba",
        "JEAN PIERRE PRUEBA EJEMPLO MINISTRO": "JEAN PIERRE PRUEBA EJEMPLO",
        "LEOPOLDO ANDRES PRUEBA EJEMPLO": "LEOPOLDO ANDRES PRUEBA EJEMPLO",
        "ADELITA INES EJEMPLO PRUEBA": "ADELITA INES EJEMPLO PRUEBA",
        "Leopoldo Andrés Prueba E.": "Leopoldo Andrés Prueba",
        "Adelita Inés Ejemplo P.": "Adelita Inés Ejemplo",
        "Jean Pierre Prueba E.": "Jean Pierre Prueba",
        "Abogado Integrante Carlos Antonio Prueba E.": "Carlos Antonio Prueba",
        "Michelle Dupont": "Michelle Dupont",
        "William Smith": "William Smith",
        "Nathalie Soto": "Nathalie Soto",
    }
    for source, expected in cases.items():
        assert _codes(source, ["NOMBRE"]) == {"NOMBRE": expected}, source
    assert _codes("ABOGADO INTEGRANTE Fecha: 22/04/2026", ["NOMBRE"]) == {}
    assert _codes("Corte Suprema, Ministerio de Hacienda, Fecha de Resolución", ["NOMBRE"]) == {}


def test_lista_de_personas_no_pierde_el_ultimo_apellido():
    for text, expected in [
        ("Rodrigo Prueba Ejemplo, Rene Ejemplo Prueba y Diego Prueba Ejemplo, con fecha",
         ["Rodrigo Prueba Ejemplo", "Rene Ejemplo Prueba", "Diego Prueba Ejemplo"]),
        ("Sr. Rene Ejemplo Prueba y Diego Prueba Ejemplo",
         ["Rene Ejemplo Prueba", "Diego Prueba Ejemplo"]),
        ("RENE EJEMPLO PRUEBA Y DIEGO PRUEBA EJEMPLO",
         ["RENE EJEMPLO PRUEBA", "DIEGO PRUEBA EJEMPLO"]),
        ("María Pérez y Juan Arriagada", ["María Pérez", "Juan Arriagada"]),
        ("José Ortega y Gasset", ["José Ortega y Gasset"]),
        ("María Pérez y Soto", ["María Pérez y Soto"]),
    ]:
        assert [text[s.start:s.end] for s in engine.detect(text, ["NOMBRE"])] == expected


def test_apellidos_chilenos_apoyan_nombres_fuera_del_catalogo():
    for text in ["Rayén Arriagada Huenchumán", "Aylén Paillalef Catrileo",
                 "Millaray Antilef Huenchuleo", "RAYEN ARRIAGADA HUENCHUMAN",
                 "Yazmín Sepúlveda Oyarzún"]:
        assert _codes(text, ["NOMBRE"]) == {"NOMBRE": text}
    for text in ["Arriagada informó los resultados", "Valencia", "Los Lagos",
                 "Sociedad Arriagada Vergara", "Colegio Arriagada Vergara",
                 "Comisión Evaluadora", "Ministerio de Hacienda"]:
        assert _codes(text, ["NOMBRE"]) == {}, text
def test_telefono_antiguo_requiere_etiqueta_y_no_confunde_fechas():
    from backend import engine
    for label in ['Teléfono:', 'tel fono', 'Celular']:
        text=label+' 09-8123456'
        found=engine.detect(text,['TELEFONO'])
        assert [text[f.start:f.end] for f in found]==['09-8123456']
    assert not engine.detect('09-8123456, 2013-2017',['TELEFONO'])
def test_apellido_de_parte_judicial_con_contexto():
    text='causa laboral caratulado "SCHNEIDER con INDUSTRIAS S.A."'
    spans=engine.detect(text,['NOMBRE'])
    assert [text[s.start:s.end] for s in spans]==['SCHNEIDER']
    assert not engine.detect('caratulado "BANKBOSTON N.A. con INDUSTRIAS S.A."',['NOMBRE'])

def test_arroba_ocr_partida_en_palabras_conserva_offsets():
    text='Correo ventas(M empresa.cl, contacto comercial'
    spans=engine.detect(text,['EMAIL'])
    assert [text[s.start:s.end] for s in spans]==['ventas(M empresa.cl']
    assert not engine.detect('suma (M) cuatro millones',['EMAIL'])

def test_nombres_en_instituciones_y_marcas_son_detecciones_validas():
    text='Universidad Diego Portales, Instituto Gabriela Mistral y vino Don Mateo.'
    spans=engine.detect(text,['NOMBRE'])
    detected=[text[s.start:s.end] for s in spans]
    assert 'Diego Portales' in detected
    assert any('Gabriela Mistral' in value for value in detected)
    assert 'Mateo' in detected
