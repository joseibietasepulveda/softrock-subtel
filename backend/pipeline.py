"""
Pipeline por etapas para revisión humana (v1.1):

  extract(src, workdir)            → extracción serializable (páginas+palabras o textos)
  detect_findings(extraction, …)   → hallazgos [dict] con ancla (page+boxes o location+start/end)
  apply_findings(src, extraction, findings, dst) → documento protegido (solo hallazgos aceptados/editados)
  verify(dst, suppressed)          → re-extrae y comprueba que nada suprimido siga presente

Los hallazgos son dicts JSON-compatibles; la Plataforma los persiste en su BD.
"""
from __future__ import annotations

import io
import json
import os
import re
import shutil
import tempfile
import unicodedata
import uuid
from pathlib import Path

from . import engine, explicit_names
from .engine import BLOCK, EngineError, _convert, _ocr_words, _signature_boxes, detect

PAGED_FORMATS = {".pdf", ".jpg", ".jpeg", ".png", ".tif", ".tiff"}
TEXT_FORMATS = {".docx", ".xlsx", ".txt", ".doc", ".xls"}
DPI = 200

TREATMENTS = {
    "redact": "Tachar",
    "anonymize": "Anonimizar",
    "pseudonymize": "Seudonimizar",
    "mask_full": "Enmascarar total",
    "mask_partial": "Enmascarar parcial",
}

ANALYSIS_LEVELS = {
    "fast": "Rápido",
    "balanced": "Equilibrado",
    "exhaustive": "Exhaustivo",
}

_ANON_LABELS = {
    "NOMBRE": "NOMBRE", "RUT": "RUT", "DIRECCION": "DIRECCIÓN",
    "EMAIL": "CORREO", "TELEFONO": "TELÉFONO", "FIRMA": "FIRMA",
    "FECHA_NACIMIENTO": "FECHA NACIMIENTO", "FECHA": "FECHA",
    "PATENTE": "PATENTE", "IP": "IP", "CUENTA_BANCARIA": "CUENTA BANCARIA",
    "PASAPORTE": "PASAPORTE", "TARJETA": "TARJETA", "WEB_REDES": "WEB/RED",
    "FOLIO": "FOLIO",
}

_PSEUDONYM_PREFIXES = {
    "NOMBRE": "Persona", "RUT": "RUT", "IP": "IP",
    "DIRECCION": "Dirección", "EMAIL": "Correo",
    "TELEFONO": "Teléfono", "FIRMA": "Firma", "FECHA_NACIMIENTO": "FechaNacimiento",
    "FECHA": "Fecha", "PATENTE": "Patente", "CUENTA_BANCARIA": "Cuenta",
    "PASAPORTE": "Pasaporte", "TARJETA": "Tarjeta", "WEB_REDES": "WebRed",
    "FOLIO": "Folio",
}


def max_pages() -> int:
    """Tope operacional de páginas por documento (0 = sin límite)."""
    return int(os.environ.get("MAX_PAGES", "500"))


def _check_pages(n: int):
    m = max_pages()
    if m and n > m:
        raise EngineError("TOO_MANY_PAGES",
                          f"El documento tiene {n} páginas y el máximo configurado es {m} (variable MAX_PAGES)")

_SCORES = {"NOMBRE": 0.7, "DIRECCION": 0.8, "FIRMA": 0.4}


def _finding(code, text, **kw):
    return {"id": uuid.uuid4().hex, "entity_code": code, "text": text,
            "score": _SCORES.get(code, 0.95), "source": "auto", "status": "accepted",
            "page": None, "boxes": None, "location": None, "start": None, "end": None, **kw}


# ---------------------------------------------------------------------------
# extract
# ---------------------------------------------------------------------------

def extract(src: Path, workdir: Path, analysis_level: str = "fast") -> dict:
    src, workdir = Path(src), Path(workdir)
    ext = src.suffix.lower()
    workdir.mkdir(parents=True, exist_ok=True)
    if ext not in engine.SUPPORTED:
        raise EngineError("UNSUPPORTED_FORMAT", f"Formato no soportado: {ext}")
    if analysis_level not in ANALYSIS_LEVELS:
        raise EngineError("INVALID_ANALYSIS_LEVEL", f"Nivel de análisis inválido: {analysis_level}")

    if ext in (".doc", ".xls"):
        conv = _convert(src, "docx" if ext == ".doc" else "xlsx", workdir)
        e = _extract_text(conv, "docx" if ext == ".doc" else "xlsx")
        e.update(format=ext[1:], converted=str(conv))
    elif ext in TEXT_FORMATS:
        e = _extract_text(src, ext[1:])
    else:
        e = _extract_paged(src, workdir, analysis_level)
    e["analysis_level"] = analysis_level
    (workdir / "extraction.json").write_text(json.dumps(e, ensure_ascii=False))
    return e


def _save_page(img, path: Path) -> None:
    """Escribe la imagen de la página de forma atómica (nunca queda un PNG a medias)."""
    path = Path(path)
    tmp = path.with_name(f".{path.name}.{uuid.uuid4().hex}.tmp")
    try:
        img.save(tmp, format="PNG")
        tmp.replace(path)
    finally:
        tmp.unlink(missing_ok=True)


def _page_has_full_scan(page) -> bool:
    """Detecta PDFs escaneados que además traen una capa OCR incompleta.

    Algunos escáneres guardan una imagen de página completa y encima una capa de
    texto parcial. Confiar sólo en esa capa deja datos visibles sin una caja que
    pueda tacharlos.
    """
    page_area = max(float(page.width) * float(page.height), 1.0)
    for image in page.images:
        try:
            image_area = abs(float(image["x1"]) - float(image["x0"])) * abs(
                float(image["bottom"]) - float(image["top"])
            )
        except (KeyError, TypeError, ValueError):
            continue
        if image_area >= page_area * 0.7:
            return True
    return False


_BALANCED_OCR_ANCHORS = re.compile(
    r"\b(?:RUT|NOMBRE|APELLIDOS?|REPRESENTANTE|RESIDENCIA|DIRECCI[ÓO]N|"
    r"CORREO|E-?MAIL|PATENTE|ROL|DISTRIBUCI[ÓO]N|FIRMA)\b",
    re.IGNORECASE,
)


def _page_needs_balanced_ocr(page, extracted_words) -> bool:
    """Selecciona páginas híbridas con mayor riesgo de capa de texto incompleta.

    Es una decisión barata y determinista basada en la propia capa PDF: estructura
    tabular, anclas de datos personales, poco texto o texto especialmente pequeño.
    No ejecuta OCR para decidir, porque eso eliminaría el ahorro del perfil.
    """
    text = " ".join(str(word.get("text", "")) for word in extracted_words)
    if _BALANCED_OCR_ANCHORS.search(text):
        return True
    try:
        shapes = len(page.lines) + len(page.rects)
    except (AttributeError, TypeError):
        shapes = 0
    if shapes >= 12 or len(extracted_words) < 35:
        return True
    heights = []
    for word in extracted_words:
        try:
            height = float(word["bottom"]) - float(word["top"])
        except (KeyError, TypeError, ValueError):
            continue
        if height > 0:
            heights.append(height)
    if heights:
        heights.sort()
        median_height = heights[len(heights) // 2]
        if median_height <= 5.5:
            return True
    return False


def _hybrid_ocr_passes(analysis_level: str, page, extracted_words) -> tuple[bool, bool]:
    """Devuelve (OCR ampliado, OCR de verificación) para una página con texto nativo."""
    if not _page_has_full_scan(page) or analysis_level == "fast":
        return False, False
    if analysis_level == "exhaustive":
        return True, True
    return _page_needs_balanced_ocr(page, extracted_words), False


def _word_box_covered(candidate, existing) -> bool:
    """Indica si una palabra OCR ya está representada por la capa de texto."""
    _, x0, y0, x1, y1 = candidate
    area = max((x1 - x0) * (y1 - y0), 1)
    cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
    for word in existing:
        _, ex0, ey0, ex1, ey1 = word
        intersection = max(0, min(x1, ex1) - max(x0, ex0)) * max(0, min(y1, ey1) - max(y0, ey0))
        if (intersection / area >= 0.45
                or ex0 - 2 <= cx <= ex1 + 2 and ey0 - 2 <= cy <= ey1 + 2):
            return True
    return False


def _sparse_ocr_words(img, scale: int = 2) -> list:
    """OCR de texto pequeño/disperso, conservando coordenadas de la imagen original."""
    import pytesseract

    enlarged = img.resize((img.width * scale, img.height * scale))
    try:
        kwargs = {"output_type": pytesseract.Output.DICT, "config": "--psm 11"}
        lang = engine._ocr_language()
        if lang:
            kwargs["lang"] = lang
        data = pytesseract.image_to_data(enlarged, **kwargs)
    finally:
        enlarged.close()
    words = []
    for i, text in enumerate(data["text"]):
        text = text.strip()
        if not text:
            continue
        x, y, width, height = (data["left"][i], data["top"][i],
                               data["width"][i], data["height"][i])
        words.append((text, x // scale, y // scale,
                      (x + width + scale - 1) // scale,
                      (y + height + scale - 1) // scale))
    return words


def _compressed_ocr_words(img) -> list:
    """OCR sobre la misma compresión JPEG usada al reconstruir un PDF protegido."""
    jpeg = io.BytesIO()
    rgb = img.convert("RGB")
    try:
        rgb.save(jpeg, format="JPEG", quality=85)
    finally:
        rgb.close()
    jpeg.seek(0)
    from PIL import Image
    with Image.open(jpeg) as compressed:
        return _ocr_words(compressed)


def _supplement_scanned_words(img, native_words, start_ratio: float = 0.0) -> list:
    """Completa la capa OCR de un escaneo sin duplicar las palabras existentes.

    Se recorre la página completa: el corpus real omite tanto tablas y pies como
    algunas patentes de la zona superior. La pasada ampliada replica lo que verá
    después el verificador sobre el PDF rasterizado.
    """
    y_offset = int(img.height * start_ratio)
    crop = img.crop((0, y_offset, img.width, img.height))
    try:
        sparse = _sparse_ocr_words(crop)
        ocr_words = [(text, x0, y0 + y_offset, x1, y1 + y_offset)
                     for text, x0, y0, x1, y1 in sparse]
    finally:
        crop.close()
    supplement = []
    for word in ocr_words:
        if not _word_box_covered(word, [*native_words, *supplement]):
            supplement.append(word)
    return supplement


def _scan_gap_regions(words, height: int, limit: int = 2) -> list[tuple[int, int]]:
    """Find short internal bands skipped by OCR, often caused by underlining.

    Use the page's typical word height and actual vertical coverage rather than
    document names, content-specific coordinates, or another full-page OCR pass.
    """
    heights=sorted(w[4]-w[2] for w in words if w[4]>w[2] and len(w[0])>=2)
    if not heights:return []
    typical=heights[len(heights)//2]
    bands=sorted((w[2],w[4]) for w in words if len(w[0])>=2 and 0<w[4]-w[2]<=typical*2.5
                 and height*.25<=w[2]<height*.9)
    merged=[]
    for top,bottom in bands:
        if merged and top<=merged[-1][1]:merged[-1][1]=max(bottom,merged[-1][1])
        else:merged.append([top,bottom])
    gaps=[(max(0,int(a[1]-typical*.35)),min(height,int(b[0]+typical*.35)))
          for a,b in zip(merged,merged[1:]) if typical*2<=b[0]-a[1]<=typical*8]
    return sorted(sorted(gaps,key=lambda r:r[1]-r[0],reverse=True)[:limit])


def _supplement_scan_gaps(img, words, analysis_level: str) -> tuple[list,int]:
    regions=_scan_gap_regions(words,img.height,4 if analysis_level=='exhaustive' else 2)
    supplement=[]
    for top,bottom in regions:
        with img.crop((0,top,img.width,bottom)) as crop:
            candidates=_sparse_ocr_words(crop,scale=1)
        for text,x0,y0,x1,y1 in candidates:
            candidate=(text,x0,y0+top,x1,y1+top)
            if not _word_box_covered(candidate,[*words,*supplement]):supplement.append(list(candidate))
    return supplement,len(regions)


def _extract_paged(src: Path, workdir: Path, analysis_level: str = "fast") -> dict:
    from PIL import Image, ImageSequence
    pages_dir = workdir / "pages"
    pages_dir.mkdir(exist_ok=True)
    ext = src.suffix.lower()
    pages, ocr_pages, hybrid_ocr_pages, verification_ocr_pages = [], [], [], []
    document_reference_seen = False

    if ext == ".pdf":
        import pdfplumber
        from pdf2image import convert_from_path
        try:
            pdf = pdfplumber.open(src)
        except Exception as exc:
            raise EngineError("CORRUPT_FILE", f"No se pudo abrir el PDF: {exc}")
        _check_pages(len(pdf.pages))
        for n, page in enumerate(pdf.pages, 1):
            # Una página a la vez: la memoria queda acotada por página, no por documento
            # (antes convert_from_path rasterizaba el PDF completo: >1 GB en 100 páginas).
            img = convert_from_path(str(src), dpi=DPI, first_page=n, last_page=n)[0]
            pw = page.extract_words(use_text_flow=True) or []
            verification_words = []
            if sum(ch.isalnum() for w in pw for ch in w["text"]) >= 30:
                words = [list(w) for w in engine.pdf_word_boxes(page, img.width, img.height, pw)]
                use_supplement, use_verification = _hybrid_ocr_passes(analysis_level, page, pw)
                if use_supplement:
                    supplement = _supplement_scanned_words(img, words)
                    if supplement:
                        words.extend([list(w) for w in supplement])
                        hybrid_ocr_pages.append(n)
                if use_verification:
                    verification_words = [list(w) for w in _compressed_ocr_words(img)]
                    verification_ocr_pages.append(n)
            else:
                ocr_pages.append(n)
                words = [list(w) for w in _ocr_words(img)]
                words = explicit_names.supplement_document_header(img, words) + words
                document_reference_seen |= explicit_names.has_document_reference(words)
                if document_reference_seen:
                    words += explicit_names.supplement_document_paragraphs(img, words)
                if analysis_level != 'fast':
                    supplement,_ = _supplement_scan_gaps(img,words,analysis_level)
                    words.extend(supplement)
            _save_page(img, pages_dir / f"p{n}.png")
            page_data = {"n": n, "image": f"pages/p{n}.png", "width": img.width,
                         "height": img.height, "words": words}
            if verification_words:
                page_data["verification_words"] = verification_words
            pages.append(page_data)
            img.close()
            page.flush_cache()
        pdf.close()
    else:
        im = Image.open(src)
        _check_pages(getattr(im, "n_frames", 1))
        for n, frame in enumerate(ImageSequence.Iterator(im), 1):
            f = frame.convert("RGB")
            _save_page(f, pages_dir / f"p{n}.png")
            ocr_pages.append(n)
            pages.append({"n": n, "image": f"pages/p{n}.png", "width": f.width, "height": f.height,
                          "words": [list(w) for w in _ocr_words(f)]})
            f.close()
    return {"kind": "paged", "format": ext[1:].replace("jpeg", "jpg"), "pages": pages,
            "ocr_pages": ocr_pages, "hybrid_ocr_pages": hybrid_ocr_pages,
            "verification_ocr_pages": verification_ocr_pages,
            "analysis_level": analysis_level}


def page_image_ok(path: Path) -> bool:
    """La imagen de la página existe y está completa (PNG con su marca final IEND)."""
    try:
        if Path(path).stat().st_size < 100:
            return False
        with Path(path).open("rb") as fh:
            if fh.read(8) != b"\x89PNG\r\n\x1a\n":
                return False
            fh.seek(-12, 2)
            return b"IEND" in fh.read(12)
    except OSError:
        return False


def render_page(src: Path, workdir: Path, n: int) -> Path | None:
    """Regenera la imagen de la página `n` con la misma resolución que `extract`.

    Permite recuperar la vista de revisión cuando la imagen de trabajo falta o
    quedó incompleta (reinicio del servicio, restauración de respaldo, escritura
    interrumpida) sin volver a procesar el documento.
    """
    src, workdir = Path(src), Path(workdir)
    if n < 1 or not src.exists():
        return None
    out = workdir / "pages" / f"p{n}.png"
    out.parent.mkdir(parents=True, exist_ok=True)
    tmp = out.with_name(f".{out.name}.{uuid.uuid4().hex}.tmp")
    ext = src.suffix.lower()
    try:
        if ext == ".pdf":
            from pdf2image import convert_from_path
            images = convert_from_path(str(src), dpi=DPI, first_page=n, last_page=n)
            if not images:
                return None
            images[0].save(tmp, format="PNG")
        elif ext in PAGED_FORMATS:
            from PIL import Image, ImageSequence
            with Image.open(src) as im:
                for i, frame in enumerate(ImageSequence.Iterator(im), 1):
                    if i == n:
                        frame.convert("RGB").save(tmp, format="PNG")
                        break
                else:
                    return None
        else:
            return None
        tmp.replace(out)  # publicación atómica: nunca se sirve un PNG a medio escribir
    except Exception:  # noqa: BLE001
        tmp.unlink(missing_ok=True)
        return None
    return out if page_image_ok(out) else None


def _extract_text(src: Path, fmt: str) -> dict:
    texts, order = {}, []

    def put(loc, t):
        texts[loc] = t
        order.append(loc)

    if fmt == "docx":
        import docx
        d = docx.Document(str(src))
        for loc, ts in engine._docx_paras(d):
            put(loc, "".join(t.text or "" for t in ts))
    elif fmt == "xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(src))
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.strip():
                        put(f"{ws.title}!{cell.coordinate}", cell.value)
    else:  # txt
        raw = Path(src).read_bytes()
        try:
            put("txt", raw.decode("utf-8"))
        except UnicodeDecodeError:
            put("txt", raw.decode("latin-1"))
    return {"kind": "text", "format": fmt, "texts": texts, "order": order}


# ---------------------------------------------------------------------------
# detect
# ---------------------------------------------------------------------------

def _detect_page_text_findings(page: dict, words: list, text_codes: list[str]) -> list[dict]:
    text, idx_at = engine._words_text_map(words)
    findings = []
    for sp in detect(text, text_codes):
        hit = sorted({i for i in idx_at[sp.start:sp.end] if i is not None})
        boxes = [[words[i][1], words[i][2], words[i][3], words[i][4]] for i in hit]
        if sp.code == "EMAIL" and boxes:
            # Algunas capas OCR declaran una caja demasiado corta aunque su
            # texto incluya el TLD; ampliar la última evita dejar ".cl" visible.
            h = max(boxes[-1][3] - boxes[-1][1], 1)
            boxes[-1][2] = min(page["width"], boxes[-1][2] + int(h * 1.8))
        findings.append(_finding(sp.code, text[sp.start:sp.end], page=page["n"], boxes=boxes))
    layer_page = {**page, "words": words}
    return _enhance_paged_findings(layer_page, findings, set(text_codes))


def _merge_alternate_findings(primary: list[dict], alternate: list[dict]) -> None:
    """Añade hallazgos de otra lectura OCR sólo donde la principal no cubre sus cajas."""
    for candidate in alternate:
        same_code_boxes = [b for finding in primary
                           if finding["entity_code"] == candidate["entity_code"]
                           for b in finding.get("boxes") or []]
        boxes = candidate.get("boxes") or []
        if boxes and all(any(_box_overlap(box, existing) for existing in same_code_boxes)
                         for box in boxes):
            continue
        primary.append(candidate)


def detect_findings(extraction: dict, codes: list[str], workdir: Path | None = None) -> list[dict]:
    findings = []
    text_codes = [c for c in codes if c != "FIRMA"]
    if extraction["kind"] == "paged":
        from PIL import Image
        for p in extraction["pages"]:
            words = p["words"]
            page_findings = _detect_page_text_findings(p, words, text_codes)
            alternate_words = p.get("verification_words") or []
            if alternate_words:
                alternate = _detect_page_text_findings(p, alternate_words, text_codes)
                _merge_alternate_findings(page_findings, alternate)
            findings.extend(page_findings)
            if "FIRMA" in codes and workdir:
                img = Image.open(Path(workdir) / p["image"])
                for b in _signature_boxes(img, [tuple(w) for w in words]):
                    findings.append(_finding("FIRMA", "(firma manuscrita)", page=p["n"], boxes=[list(b)]))
    else:
        for loc in extraction["order"]:
            for sp in detect(extraction["texts"][loc], text_codes):
                findings.append(_finding(sp.code, extraction["texts"][loc][sp.start:sp.end],
                                         location=loc, start=sp.start, end=sp.end))
    if extraction["kind"] == "text" and extraction["format"] in ("xlsx", "xls"):
        findings += _columnar_xlsx(extraction, codes, findings)
    findings += _propagate(findings, extraction)
    findings += _propagate_fuzzy_three_part_names(findings, extraction)
    _remove_redundant_name_findings(findings)
    if "NOMBRE" in codes:
        for name in explicit_names.exact_name_findings(extraction):
            page_findings = [f for f in findings if f.get("page") == name["page"]]
            other_pages = [f for f in findings if f.get("page") != name["page"]]
            _merge_context_name(page_findings, _finding("NOMBRE", **name))
            findings = other_pages + page_findings
    return findings


_TITLE_WORDS = {"sr", "sra", "srta", "senor", "senora", "don", "dona", "ria"}
_NAME_ROLE_WORDS = engine.WHITELIST | _TITLE_WORDS | {
    "abogado", "abogada", "profesional", "somete", "comunica", "rut", "cedula",
    "residencia", "oferente", "firma", "depto", "oficina", "piso", "bodega",
    "patente", "motor", "chasis", "modelo", "marca",
    "registros", "publicista", "periodista", "ingeniero", "ingeniera",
    "licenciado", "licenciada", "contacto", "formacion", "academica", "academico",
}


def _box_overlap(a, b) -> bool:
    return a[0] < b[2] and a[2] > b[0] and a[1] < b[3] and a[3] > b[1]


def _merge_context_name(findings: list[dict], candidate: dict) -> None:
    overlapping = [f for f in findings if f["entity_code"] == "NOMBRE" and
                   any(_box_overlap(a, b) for a in f.get("boxes") or [] for b in candidate["boxes"])]
    for finding in overlapping:
        findings.remove(finding)
    findings.append(candidate)


def _title_context_names(page: dict) -> list[dict]:
    """Lee nombres después de Sr./Sra./Srta. usando la geometría de la línea.

    Es intencionalmente tolerante al OCR de listados pequeños ("$ria", iniciales y
    nombres mal leídos): el título aporta el contexto que el diccionario no tiene.
    """
    words = page["words"]
    out = []
    for anchor in words:
        title = _norm_soft(anchor[0]).replace(" ", "")
        if title not in _TITLE_WORDS:
            continue
        height = max(anchor[4] - anchor[2], 1)
        line = sorted(
            (w for w in words if w is not anchor and engine._same_text_line(anchor, w)
             and anchor[3] - height * 0.35 <= w[1] <= anchor[1] + max(680, height * 24)),
            key=lambda w: w[1],
        )
        # Dos excepciones explícitas; no inferir una regla para «Sr + cargo».
        following = " ".join(w[0] for w in line)
        # «Señor y dueño» es la expresión del escrito, no un nombre de persona.
        if title == "senor" and _norm_soft(following).split()[:2] == ["y", "dueno"]:
            continue
        if title == "sr" and re.match(r"(?:\(a\)\s*)?(?:ministro|secretario)\b",
                                       following, re.IGNORECASE):
            continue
        picked, last_x, digit_break = [], anchor[3], False
        for word in line:
            token = _norm_soft(word[0]).replace(" ", "")
            if word[1] - last_x > height * 4 or token in _NAME_ROLE_WORDS or not token:
                break
            if token.isdigit():
                digit_break = True
                break
            if not any(ch.isalpha() for ch in token):
                if len(token) <= 2:  # inicial OCR como "J"
                    picked.append(word)
                    last_x = word[3]
                    continue
                break
            picked.append(word)
            last_x = word[3]
            if len(token) > 2 and re.search(r"[,;:]$", word[0]):
                break
            if len(picked) == 6:
                break
        if digit_break or len(picked) < 2 or _norm_soft(picked[0][0]).replace(" ", "") in _NAME_ROLE_WORDS:
            continue
        text = " ".join(w[0].strip(" ,.;:") for w in picked).strip()
        if text:
            out.append(_finding("NOMBRE", text, page=page["n"],
                                boxes=[list(w[1:]) for w in picked], score=0.88))
    return out


def _role_context_names(page: dict) -> list[dict]:
    """Recover an uncommon name immediately before a personal professional role.

    Native PDFs may return a surname as adjacent fragments (M + üller), or put
    a displaced accent between those fragments in text-flow order. Geometry
    gives the actual line order without requiring a list of literal names.
    """
    words = page['words']
    out = []
    roles = {'juez', 'jueza', 'notario', 'notaria'}
    for anchor in words:
        if _norm_soft(anchor[0]) not in roles:
            continue
        height = max(anchor[4] - anchor[2], 1)
        line = sorted((w for w in words if w[3] <= anchor[1]
                       and anchor[1] - w[3] < height * 20
                       and engine._same_text_line(anchor, w)), key=lambda w: w[1])
        joined = []
        for word in line:
            text = word[0].strip(' ,.;:')
            if (joined and text and text[0].islower()
                    and len(joined[-1]['text']) == 1
                    and joined[-1]['text'].isupper()
                    and abs(word[1] - joined[-1]['boxes'][-1][2]) <= height * 0.1):
                joined[-1]['text'] += text
                joined[-1]['boxes'].append(list(word[1:]))
            else:
                joined.append({'text': text, 'boxes': [list(word[1:])]})
        picked = []
        for item in reversed(joined):
            text = item['text']
            if (not re.fullmatch(engine._CAP, text)
                    or _norm_soft(text) in _NAME_ROLE_WORDS | {'civil', 'penal', 'laboral', 'titular'}):
                break
            if picked:
                next_left = min(b[0] for b in picked[-1]['boxes'])
                if next_left - max(b[2] for b in item['boxes']) > height * 2:
                    break
            picked.append(item)
            if len(picked) == 4:
                break
        if len(picked) >= 2:
            picked.reverse()
            out.append(_finding('NOMBRE', ' '.join(item['text'] for item in picked),
                                page=page['n'], boxes=[b for item in picked for b in item['boxes']],
                                score=0.85))
    return out


def _representative_names(page: dict, findings: list[dict]) -> list[dict]:
    """Recupera representantes de tablas a partir del encabezado y su RUT natural.

    La banda vertical entre dos RUT consecutivos permite unir apellidos partidos en
    dos líneas sin absorber la columna contigua de residencia.
    """
    words = page["words"]
    rep = next((w for w in words if _norm_soft(w[0]) == "representante"), None)
    residence = next((w for w in words if _norm_soft(w[0]) == "residencia"), None)
    ruts = sorted((f for f in findings if f["entity_code"] == "RUT" and f.get("boxes")),
                  key=lambda f: min(b[1] for b in f["boxes"]))
    if not rep or not residence or not ruts:
        return []
    left, right = rep[1] - 55, residence[1] - 40
    previous_bottom = max([rep[4], *(w[4] for w in words
                                     if _norm_soft(w[0]) == "legal" and w[1] < right)])
    out = []
    for rut in ruts:
        rut_top = min(b[1] for b in rut["boxes"])
        candidates = []
        for w in words:
            center = (w[1] + w[3]) / 2
            token = _norm_soft(w[0]).replace(" ", "")
            if not (left <= center <= right and previous_bottom - 4 <= w[2]
                    and w[2] <= rut_top + 6):
                continue
            if token in _NAME_ROLE_WORDS or any(ch.isdigit() for ch in token) or not any(ch.isalpha() for ch in token):
                continue
            candidates.append(w)
        previous_bottom = max(b[3] for b in rut["boxes"])
        candidates.sort(key=lambda w: (w[2], w[1]))
        if len(candidates) < 2:
            continue
        text = " ".join(w[0].strip(" ,.;:") for w in candidates)
        out.append(_finding("NOMBRE", text, page=page["n"],
                            boxes=[list(w[1:]) for w in candidates], score=0.9))
    return out


def _extend_wrapped_names(page: dict, findings: list[dict]) -> None:
    from statistics import median
    words = page["words"]
    heights = [w[4] - w[2] for w in words if w[4] > w[2]]
    normal_height = median(heights) if heights else 1
    for finding in [f for f in findings if f["entity_code"] == "NOMBRE" and f.get("boxes")
                    and f.get("score", 1) <= 0.7]:
        boxes = finding["boxes"]
        bottom = max(b[3] for b in boxes)
        left, right = min(b[0] for b in boxes), max(b[2] for b in boxes)
        height = max(max(b[3] - b[1] for b in boxes), 1)
        last_words = [w for w in words if list(w[1:]) in boxes]
        if last_words:
            last_word = max(last_words, key=lambda w: (w[2], w[1]))
            if last_word[0].rstrip().endswith((",", ";", ".", ":")):
                continue  # el nombre ya terminó antes del salto de línea
            # A surname at the start of a wrapped list may lie far to the left
            # of the first name. Require both consecutive extraction order and
            # an explicit list boundary after it, rather than joining columns.
            index = next((i for i, word in enumerate(words) if word is last_word), None)
            if index is not None and index + 2 < len(words):
                surname, following = words[index + 1:index + 3]
                value = surname[0].strip(' ,.;:')
                list_boundary = (surname[0].endswith((',', ';'))
                                 or _norm_soft(following[0]) == 'y')
                if (0 <= surname[2] - bottom <= height * 0.8 and list_boundary
                        and re.fullmatch(engine._CAP, value)
                        and _norm_soft(value) not in _NAME_ROLE_WORDS
                        and not engine._same_text_line(last_word, surname)):
                    finding['text'] += '\n' + value
                    finding['boxes'].append(list(surname[1:]))
                    continue
                # Two surnames can wrap together after a first name at the
                # far right margin. Their final comma establishes the boundary.
                pair=words[index+1:index+3]
                if (pair[-1][0].endswith((',', ';'))
                        and all(re.fullmatch(engine._CAP,w[0].strip(' ,.;:'))
                                and _norm_soft(w[0].strip(' ,.;:')) not in _NAME_ROLE_WORDS for w in pair)
                        and 0<=min(w[2] for w in pair)-bottom<=height*1.5
                        and engine._same_text_line(*pair)
                        and not engine._same_text_line(last_word,pair[0])):
                    finding['text']+='\n'+' '.join(w[0].strip(' ,.;:') for w in pair)
                    finding['boxes'].extend(list(w[1:]) for w in pair)
                    continue
        next_words = []
        for word in words:
            box = list(word[1:])
            token = _norm_soft(word[0]).replace(" ", "")
            horizontal_overlap = min(right, box[2]) - max(left, box[0])
            if (0 <= box[1] - bottom <= height * 0.65 and horizontal_overlap > 0
                    and token not in _NAME_ROLE_WORDS and any(ch.isalpha() for ch in token)
                    and not any(box == old for old in boxes)):
                next_words.append(word)
        if not next_words:
            continue
        top = min(w[2] for w in next_words)
        next_words = sorted((w for w in next_words if abs(w[2] - top) <= height * 0.35),
                            key=lambda w: w[1])[:2]
        prominent_heading = (height >= normal_height * 1.35
                             and bottom <= page["height"] * 0.4)
        if abs(next_words[0][1] - left) <= height * 0.45 and not prominent_heading:
            continue  # nueva fila alineada con la anterior, no apellido envuelto
        surname_words = []
        for word in next_words:
            value = word[0].strip(" ,.;:")
            if (engine._strip(value) in _NAME_ROLE_WORDS | {"de", "del", "la", "las", "los"}
                    or not (re.fullmatch(engine._CAP, value)
                            or engine._strip(value) in engine.CHILEAN_SURNAMES)):
                break
            surname_words.append(word)
        if surname_words:
            finding["text"] += "\n" + " ".join(w[0].strip(" ,.;:") for w in surname_words)
            finding["boxes"].extend([list(w[1:]) for w in surname_words])


def _labeled_ruts(page: dict, findings: list[dict]) -> None:
    """En escaneos tabulares acepta un RUT natural rotulado aunque el OCR altere el DV."""
    words = page["words"]
    existing = [b for f in findings if f["entity_code"] == "RUT" for b in f.get("boxes") or []]
    loose = re.compile(r"(?<!\d)\d{1,2}(?:\s*\.\s*\d{3}){2}\s*-\s*[\dkK](?!\d)")
    for anchor in words:
        if _norm_soft(anchor[0]) != "rut":
            continue
        height = max(anchor[4] - anchor[2], 1)
        following = sorted((w for w in words if engine._same_text_line(anchor, w)
                            and w[1] >= anchor[1] and w[3] > anchor[3]
                            and w[1] - anchor[3] <= height * 6), key=lambda w: w[1])[:2]
        if not following:
            continue
        raw = " ".join(w[0] for w in following)
        match = loose.search(raw)
        if not match:
            continue
        number = int(re.sub(r"\D", "", match.group().rsplit("-", 1)[0]))
        boxes = [list(w[1:]) for w in following]
        if number >= 50_000_000 or all(any(_box_overlap(box, old) for old in existing) for box in boxes):
            continue
        findings.append(_finding("RUT", match.group(), page=page["n"], boxes=boxes, score=0.82))
        existing.extend(boxes)


def _remove_split_company_names(page: dict, findings: list[dict]) -> None:
    words = page["words"]
    for finding in list(findings):
        if finding["entity_code"] != "NOMBRE" or not finding.get("boxes"):
            continue
        first_box = finding["boxes"][0]
        first_word = next((w for w in words if list(w[1:]) == first_box), None)
        if not first_word:
            continue
        previous = [w for w in words if w[4] <= first_word[2] + 3 and
                    0 <= first_word[2] - w[4] <= max(first_word[4] - first_word[2], 1) * 1.5]
        first = _norm_soft(first_word[0]).replace(" ", "")
        if any(_norm_soft(w[0]).replace(" ", "") + first == "comercializadora" for w in previous):
            findings.remove(finding)


def _remove_invalid_names(findings: list[dict]) -> None:
    """Un nombre de persona automático nunca contiene números ni vocabulario vehicular."""
    for finding in list(findings):
        if finding["entity_code"] != "NOMBRE":
            continue
        tokens = set(_norm_soft(finding.get("text") or "").split())
        if any(token.isdigit() for token in tokens) or tokens & {
            "patente", "motor", "chasis", "modelo", "marca", "interno",
        }:
            findings.remove(finding)
            continue
        if len(tokens) == 1 and finding.get("boxes"):
            if any(other is not finding and other["entity_code"] == "NOMBRE"
                   and len(_norm_soft(other.get("text") or "").split()) >= 2
                   and any(_box_overlap(a, b) for a in finding["boxes"]
                           for b in other.get("boxes") or [])
                   for other in findings):
                findings.remove(finding)


def _remove_redundant_name_findings(findings: list[dict]) -> None:
    """Retira fragmentos de un solo nombre ya cubiertos por un nombre completo."""
    for finding in list(findings):
        if finding["entity_code"] != "NOMBRE" or not finding.get("boxes"):
            continue
        if len(_norm_soft(finding.get("text") or "").split()) != 1:
            continue
        if any(other is not finding and other["entity_code"] == "NOMBRE"
               and other.get("page") == finding.get("page") and other.get("boxes")
               and len(_norm_soft(other.get("text") or "").split()) >= 2
               and any(_box_overlap(a, b) for a in finding["boxes"] for b in other["boxes"])
               for other in findings):
            findings.remove(finding)


def _residence_addresses(page: dict) -> list[dict]:
    words = page["words"]
    header = next((w for w in words if _norm_soft(w[0]) == "residencia"), None)
    if not header:
        return []
    left = header[1] - 50
    candidates = [w for w in words if (w[1] + w[3]) / 2 >= left and w[2] > header[4]]
    lines: list[list] = []
    for word in sorted(candidates, key=lambda w: (w[2], w[1])):
        line = next((line for line in lines if engine._same_text_line(line[0], word)), None)
        if line is None:
            lines.append([word])
        else:
            line.append(word)
    out = []
    for line in lines:
        line.sort(key=lambda w: w[1])
        if not any(any(ch.isdigit() for ch in w[0]) for w in line):
            continue
        first_alpha = next((i for i, w in enumerate(line) if any(ch.isalpha() for ch in w[0])), None)
        last_digit = max((i for i, w in enumerate(line) if any(ch.isdigit() for ch in w[0])), default=-1)
        if first_alpha is None or last_digit < first_alpha:
            continue
        selected = line[first_alpha:last_digit + 1]
        lead = _norm_soft(selected[0][0]).replace(" ", "")
        if lead in {"rut", "monto", "articulo", "pagina", "n"}:
            continue
        text = " ".join(w[0].strip(" ,.;:") for w in selected)
        out.append(_finding("DIRECCION", text, page=page["n"],
                            boxes=[list(w[1:]) for w in selected], score=0.88))
    return out


def _enhance_paged_findings(page: dict, findings: list[dict], codes: set[str]) -> list[dict]:
    if "RUT" in codes:
        _labeled_ruts(page, findings)
    if "NOMBRE" in codes:
        for candidate in _title_context_names(page):
            _merge_context_name(findings, candidate)
        for candidate in _role_context_names(page):
            _merge_context_name(findings, candidate)
        for candidate in _representative_names(page, findings):
            _merge_context_name(findings, candidate)
        _extend_wrapped_names(page, findings)
        _remove_split_company_names(page, findings)
        _remove_invalid_names(findings)
    if "DIRECCION" in codes:
        existing_boxes = [b for f in findings if f["entity_code"] == "DIRECCION" for b in f.get("boxes") or []]
        for candidate in _residence_addresses(page):
            if not all(any(_box_overlap(box, old) for old in existing_boxes) for box in candidate["boxes"]):
                findings.append(candidate)
                existing_boxes.extend(candidate["boxes"])
        for candidate in _wrapped_street_addresses(page):
            if not all(any(_box_overlap(box, old) for old in existing_boxes) for box in candidate['boxes']):
                findings.append(candidate)
                existing_boxes.extend(candidate['boxes'])
    return findings


def _wrapped_street_addresses(page: dict) -> list[dict]:
    """A street name at the end of a line may put its N° and number on the next.

    Require the explicit number marker and nearby physical lines; never allow
    a following year or arbitrary paragraph number to complete an address.
    """
    words=page['words'];out=[]
    for index,anchor in enumerate(words):
        if _norm_soft(anchor[0]).rstrip('.') not in {'calle','avenida','avda','av','pasaje','pje','camino'}:
            continue
        selected=[anchor];height=max(anchor[4]-anchor[2],1)
        for word in words[index+1:index+6]:
            if not engine._same_text_line(anchor,word):break
            if word[1]<selected[-1][3] or word[1]-selected[-1][3]>height*3:break
            if not re.fullmatch(engine._CAP,word[0].strip(' ,.;:')):break
            selected.append(word)
        if len(selected)<2:continue
        next_index=index+len(selected)
        if next_index+1>=len(words):continue
        marker,number=words[next_index:next_index+2]
        if not re.fullmatch(r'N\.?[°ºo.]',marker[0],re.IGNORECASE):continue
        if not re.fullmatch(r'\d{1,5}[,.;]?',number[0]):continue
        if not (0<=marker[2]-selected[-1][4]<=height*1.8 and marker[1]<selected[-1][1]
                and engine._same_text_line(marker,number) and 0<=number[1]-marker[3]<=height*3):continue
        selected.extend([marker,number])
        unit_words=words[next_index+2:next_index+4]
        if (len(unit_words)==2 and _norm_soft(unit_words[0][0]).rstrip('.') in
                {'oficina','of','departamento','depto','piso','casa','local','block'}
                and re.fullmatch(r'[A-Za-z0-9-]{1,6}[,.;]?',unit_words[1][0])
                and all(engine._same_text_line(number,w) for w in unit_words)
                and 0<=unit_words[0][1]-number[3]<=height*3):
            selected.extend(unit_words)
        out.append(_finding('DIRECCION',' '.join(w[0].strip(' ,.;:') for w in selected),
                            page=page['n'],boxes=[list(w[1:]) for w in selected],score=.88))
    return out


# Encabezado de columna → tipo de dato (planillas: el valor va "pelado" bajo el título)
_HEADER_MAP = [
    ("FOLIO", {"folio", "expediente", "ingreso"}),
    ("RUT", {"rut", "run"}),
    ("EMAIL", {"correo", "email", "e-mail"}),
    ("TELEFONO", {"telefono", "fono", "celular"}),
    ("DIRECCION", {"domicilio", "direccion"}),
    ("NOMBRE", {"nombre", "reclamante", "solicitante", "denunciante", "funcionario",
                "funcionaria", "trabajador", "interesado", "apoderado", "representante",
                "apellido", "apellidos"}),
    ("CUENTA_BANCARIA", {"cuenta"}),
    ("PATENTE", {"patente"}),
    ("PASAPORTE", {"pasaporte"}),
    ("IP", {"ip"}),
    ("FECHA_NACIMIENTO", {"nacimiento"}),
    ("TARJETA", {"tarjeta"}),
    ("FECHA", {"fecha"}),
]
_CELL_RE = re.compile(r"^(.*)!([A-Z]{1,3})(\d+)$")


def _columnar_xlsx(extraction: dict, codes: list[str], findings: list[dict]) -> list[dict]:
    """FR-30 para planillas: si el encabezado de la columna nombra un dato sensible
    ("Folio", "RUT", "Reclamante"…), todas las celdas bajo él son ese dato, aunque
    el valor no tenga contexto propio."""
    cols: dict[tuple[str, str], list[tuple[int, str]]] = {}
    for loc in extraction["order"]:
        m = _CELL_RE.match(loc)
        if m:
            cols.setdefault((m.group(1), m.group(2)), []).append((int(m.group(3)), loc))
    existing_by_loc: dict[str, list[dict]] = {}
    for finding in findings:
        if finding.get("location"):
            existing_by_loc.setdefault(finding["location"], []).append(finding)
    replaced_locs: set[str] = set()
    extra: list[dict] = []
    for (_sheet, _col), cells in cols.items():
        cells.sort()
        header_code, header_row = None, None
        for row, loc in cells:
            words = set(_norm_soft(extraction["texts"][loc]).split())
            code = next((c for c, kws in _HEADER_MAP if c in codes and words & kws), None)
            if code and len(extraction["texts"][loc]) <= 40:
                header_code, header_row = code, row
                break
        if not header_code:
            continue
        for row, loc in cells:
            if row <= header_row:
                continue
            v = extraction["texts"][loc].strip()
            if (len(v) < 3 and header_code != "NOMBRE") or v in ("-", "—", "N/A", "n/a", "s/i", "S/I"):
                continue
            i = extraction["texts"][loc].index(v[0])
            end = i + len(v)
            existing = existing_by_loc.get(loc, [])
            if len(existing) == 1 and existing[0]["entity_code"] == header_code \
                    and existing[0]["start"] == i and existing[0]["end"] == end:
                continue
            # El encabezado define el tipo de toda la celda. Un hallazgo genérico
            # parcial no debe impedir que se proteja el resto del apellido.
            replaced_locs.add(loc)
            extra.append(_finding(header_code, v, location=loc, start=i, end=i + len(v), score=0.85))
    if replaced_locs:
        findings[:] = [f for f in findings if f.get("location") not in replaced_locs]
    return extra


def _propagate(findings: list[dict], extraction: dict) -> list[dict]:
    """Repite cada valor detectado sobre TODAS sus ocurrencias literales.

    Un folio detectado junto a la palabra "folio" puede aparecer pelado en otra
    celda o página; sin esto, el verificador (con razón) bloqueaba el documento.
    """
    values: dict[str, tuple[str, str]] = {}
    for f in findings:
        t = (f["text"] or "").strip()
        if len(t) >= 6 and f["entity_code"] != "FIRMA" and not t.startswith("("):
            values.setdefault(t.lower(), (t, f["entity_code"]))
    if not values:
        return []

    def whole_value(text: str, start: int, end: int) -> bool:
        before = text[start - 1] if start else ""
        after = text[end] if end < len(text) else ""
        # En nombres, el llamador también usa esta frontera: los separadores de
        # correo no cuentan como límite para no propagar "Veronica" en un e-mail.
        return (not before.isalnum() and before not in "@._+-"
                and not after.isalnum() and after not in "@._+-")

    extra: list[dict] = []
    if extraction["kind"] == "text":
        by_loc: dict[str, list] = {}
        for f in findings:
            if f.get("location") is not None:
                by_loc.setdefault(f["location"], []).append(f)
        # texto completo unido con \n + índice de offsets → una búsqueda C por valor
        order = extraction["order"]
        offsets, big_parts, pos = [], [], 0
        for loc in order:
            offsets.append(pos)
            big_parts.append(extraction["texts"][loc].lower())
            pos += len(big_parts[-1]) + 1
        big = "\n".join(big_parts)
        import bisect
        for key, (_orig, code) in values.items():
            start = 0
            while (i := big.find(key, start)) != -1:
                start = i + 1
                if "\n" in big[i:i + len(key)]:
                    continue  # cruza de una pieza a otra: no es una ocurrencia real
                k = bisect.bisect_right(offsets, i) - 1
                loc = order[k]
                li = i - offsets[k]
                lj = li + len(key)
                if code == "NOMBRE" and not whole_value(extraction["texts"][loc], li, lj):
                    continue
                if not any(g["start"] < lj and g["end"] > li for g in by_loc.get(loc, [])):
                    nf = _finding(code, extraction["texts"][loc][li:lj], location=loc, start=li, end=lj, score=0.9)
                    extra.append(nf)
                    by_loc.setdefault(loc, []).append(nf)
    else:
        for p in extraction["pages"]:
            words = p["words"]
            text, idx_at = engine._words_text_map(words)
            low = text.lower()
            # cajas ya tachadas en esta página; una ocurrencia sólo se agrega si sus
            # palabras no están cubiertas (el mismo valor puede repetirse en la página)
            page_boxes = [b for f in findings if f["page"] == p["n"] and f["boxes"] for b in f["boxes"]]
            for key, (_orig, code) in values.items():
                start = 0
                while (i := low.find(key, start)) != -1:
                    start = i + 1
                    end = i + len(key)
                    if "\n" in low[i:end] and code != "NOMBRE":
                        continue
                    if code == "NOMBRE" and not whole_value(low, i, end):
                        continue
                    hit = sorted({k for k in idx_at[i:end] if k is not None})
                    if not hit:
                        continue
                    boxes = [[words[k][1], words[k][2], words[k][3], words[k][4]] for k in hit]
                    covered = all(any(bx0 <= (x0 + x1) / 2 <= bx1 and by0 <= (y0 + y1) / 2 <= by1
                                      for bx0, by0, bx1, by1 in page_boxes)
                                  for x0, y0, x1, y1 in boxes)
                    if covered:
                        continue
                    extra.append(_finding(code, text[i:end], page=p["n"], score=0.9, boxes=boxes))
                    page_boxes += boxes

            # El mismo nombre puede reaparecer sin espacios por OCR o por una celda
            # angosta (PAULA SAHR / HENRIQUEZ -> PAULASAHR\nHENRIQUEZ).
            compact_chars, compact_at = [], []
            for original_i, char in enumerate(text):
                normalized = "".join(c for c in unicodedata.normalize("NFD", char)
                                     if unicodedata.category(c) != "Mn").lower()
                for normalized_char in normalized:
                    if normalized_char.isalnum():
                        compact_chars.append(normalized_char)
                        compact_at.append(original_i)
            compact = "".join(compact_chars)
            compact_names = {(_norm(orig), orig) for _key, (orig, code) in values.items()
                             if code == "NOMBRE" and len(_norm(orig)) >= 6}
            for compact_key, _orig in compact_names:
                start = 0
                while (ci := compact.find(compact_key, start)) != -1:
                    start = ci + 1
                    cj = ci + len(compact_key)
                    oi, oj = compact_at[ci], compact_at[cj - 1] + 1
                    if not whole_value(text, oi, oj):
                        continue
                    hit = sorted({k for k in idx_at[oi:oj] if k is not None})
                    if not hit:
                        continue
                    boxes = [[words[k][1], words[k][2], words[k][3], words[k][4]] for k in hit]
                    covered = all(any(bx0 <= (x0 + x1) / 2 <= bx1 and by0 <= (y0 + y1) / 2 <= by1
                                      for bx0, by0, bx1, by1 in page_boxes)
                                  for x0, y0, x1, y1 in boxes)
                    if covered:
                        continue
                    extra.append(_finding("NOMBRE", text[oi:oj], page=p["n"], score=0.9, boxes=boxes))
                    page_boxes += boxes

                # Segunda pasada geométrica: omite palabras de otra columna que el
                # orden interno del PDF intercala entre dos líneas del mismo nombre.
                for first in words:
                    first_key = _norm(first[0])
                    if not first_key or first_key == compact_key or not compact_key.startswith(first_key):
                        continue
                    selected, assembled, current = [first], first_key, first
                    while assembled != compact_key:
                        height = max(current[4] - current[2], 1)
                        options = []
                        for candidate in words:
                            if candidate in selected:
                                continue
                            token = _norm(candidate[0])
                            if not token or not compact_key.startswith(assembled + token):
                                continue
                            same_line = engine._same_text_line(current, candidate)
                            horizontal_next = same_line and -height <= candidate[1] - current[3] <= height * 3
                            overlap = min(current[3], candidate[3]) - max(current[1], candidate[1])
                            wrapped_next = (0 <= candidate[2] - current[4] <= height * 0.8 and overlap > 0)
                            if horizontal_next or wrapped_next:
                                options.append(candidate)
                        if not options:
                            break
                        current = min(options, key=lambda w: (abs(w[2] - current[2]), w[1]))
                        selected.append(current)
                        assembled += _norm(current[0])
                    if assembled != compact_key:
                        continue
                    boxes = [list(w[1:]) for w in selected]
                    covered = all(any(bx0 <= (x0 + x1) / 2 <= bx1 and by0 <= (y0 + y1) / 2 <= by1
                                      for bx0, by0, bx1, by1 in page_boxes)
                                  for x0, y0, x1, y1 in boxes)
                    if covered:
                        continue
                    extra.append(_finding("NOMBRE", "\n".join(w[0] for w in selected), page=p["n"],
                                          score=0.9, boxes=boxes))
                    page_boxes += boxes
    return extra


def _propagate_fuzzy_three_part_names(findings: list[dict], extraction: dict) -> list[dict]:
    """Propaga `MIG[UEL] ARAVENA ANGULO` cuando OCR abrevia el primer nombre.

    Sólo se activa para un nombre completo de tres partes ya detectado en otra
    zona y exige coincidencia exacta de ambos apellidos en la misma línea.
    """
    if extraction.get("kind") != "paged":
        return []
    known_pages: dict[tuple[str, str, str], set[int]] = {}
    for finding in findings:
        if finding["entity_code"] != "NOMBRE":
            continue
        tokens = _norm_soft(finding.get("text") or "").split()
        if len(tokens) == 3 and all(len(token) >= 3 for token in tokens):
            known_pages.setdefault(tuple(tokens), set()).add(finding.get("page") or 0)
    known = {tokens for tokens, pages in known_pages.items()
             if len(pages) >= 2 and tokens[0] in engine.FIRST_NAMES}
    extra = []
    for page in extraction["pages"]:
        existing = [b for finding in findings if finding.get("page") == page["n"]
                    and finding["entity_code"] == "NOMBRE" for b in finding.get("boxes") or []]
        for words in (page.get("words") or [], page.get("verification_words") or []):
            for first_name, surname1, surname2 in known:
                for middle in words:
                    if _norm_soft(middle[0]).replace(" ", "") != surname1:
                        continue
                    height = max(middle[4] - middle[2], 1)
                    before = [word for word in words if word[3] <= middle[1]
                              and engine._same_text_line(word, middle)
                              and 0 <= middle[1] - word[3] <= height * 4
                              and len(_norm_soft(word[0]).replace(" ", "")) >= 3
                              and first_name.startswith(_norm_soft(word[0]).replace(" ", ""))]
                    after = [word for word in words if word[1] >= middle[3]
                             and engine._same_text_line(word, middle)
                             and 0 <= word[1] - middle[3] <= height * 4
                             and _norm_soft(word[0]).replace(" ", "") == surname2]
                    if not before or not after:
                        continue
                    selected = [max(before, key=lambda word: word[3]), middle,
                                min(after, key=lambda word: word[1])]
                    boxes = [list(word[1:]) for word in selected]
                    boxes[0][0] = max(0, boxes[0][0] - height * 6)
                    if all(any(_box_overlap(box, old) for old in existing) for box in boxes):
                        continue
                    extra.append(_finding("NOMBRE", " ".join(word[0] for word in selected),
                                          page=page["n"], boxes=boxes, score=0.82))
                    existing.extend(boxes)
                # En pies timbrados el OCR puede perder todos los componentes salvo
                # el último apellido. Como el nombre completo ya existe en otra zona,
                # se cubre la banda izquierda inmediata del apellido repetido.
                for tail in words:
                    if (_norm_soft(tail[0]).replace(" ", "") != surname2
                            or tail[2] < page["height"] * 0.75):
                        continue
                    fragments = [word for word in words if word is not tail
                                 and engine._same_text_line(word, tail)
                                 and word[3] <= tail[1]
                                 and 0 <= tail[1] - word[3] <= max(tail[4] - tail[2], 1) * 10]
                    if not any(len(token := _norm_soft(word[0]).replace(" ", "")) >= 3
                               and (first_name.startswith(token) or surname1.startswith(token))
                               for word in fragments):
                        continue
                    if any(_box_overlap(list(tail[1:]), old) for old in existing):
                        continue
                    word_h = max(tail[4] - tail[2], 1)
                    box = [max(0, tail[1] - word_h * 24), max(0, tail[2] - word_h),
                           tail[3], min(page["height"], tail[4] + word_h)]
                    extra.append(_finding("NOMBRE", " ".join((first_name, surname1, surname2)),
                                          page=page["n"], boxes=[box], score=0.78))
                    existing.append(box)
    return extra


# ---------------------------------------------------------------------------
# apply
# ---------------------------------------------------------------------------

def _redact_span(text: str, start: int, end: int) -> str:
    return text[:start] + "".join(BLOCK if not c.isspace() else c for c in text[start:end]) + text[end:]


def _pseudonym_key(finding: dict) -> str:
    return f"{finding['entity_code']}:{_norm(finding.get('text') or '')}"


def pseudonym_key(finding: dict) -> str:
    """Clave normalizada usada por la plataforma para resolver alias persistentes."""
    return _pseudonym_key(finding)


def pseudonym_prefix(entity_code: str) -> str:
    return _PSEUDONYM_PREFIXES.get(entity_code, entity_code.title())


def build_pseudonym_map(findings: list[dict]) -> dict[str, str]:
    """Construye alias estables y deterministas para un documento o lote completo."""
    grouped: dict[str, set[str]] = {}
    for finding in findings:
        if finding.get("status") not in ("accepted", "edited") or not finding.get("text"):
            continue
        grouped.setdefault(finding["entity_code"], set()).add(_norm(finding["text"]))
    result: dict[str, str] = {}
    for code, values in grouped.items():
        prefix = pseudonym_prefix(code)
        for number, value in enumerate(sorted(values), 1):
            result[f"{code}:{value}"] = f"{prefix}-{number:02d}"
    return result


def replacement_for(finding: dict, treatment: str,
                    pseudonym_map: dict[str, str] | None = None) -> str:
    """Devuelve el sustituto visible correspondiente al tratamiento elegido."""
    if treatment not in TREATMENTS:
        raise ValueError(f"Tratamiento inválido: {treatment}")
    original = finding.get("text") or ""
    code = finding["entity_code"]
    if treatment == "redact":
        return "".join(BLOCK if not c.isspace() else c for c in original)
    if treatment == "anonymize":
        return f"[{_ANON_LABELS.get(code, code)}]"
    if treatment == "pseudonymize":
        mapping = pseudonym_map or build_pseudonym_map([finding])
        return mapping.get(_pseudonym_key(finding), f"{pseudonym_prefix(code)}-01")
    if treatment == "mask_full":
        return "********"
    if code == "FIRMA":
        return "***"
    # Conserva los tres primeros caracteres alfanuméricos y la puntuación del
    # original: 12.345.678-5 → 12.3**.***-*.
    # A partial mask must still hide at least one character for short values
    # such as a three-letter first name in a spreadsheet column.
    visible = min(3, max(0, sum(char.isalnum() for char in original) - 1))
    seen = 0
    chars = []
    for char in original:
        if char.isalnum():
            seen += 1
            chars.append(char if seen <= visible else "*")
        else:
            chars.append(char)
    return "".join(chars)


def _replace_span(text: str, finding: dict, treatment: str,
                  pseudonym_map: dict[str, str] | None) -> str:
    start, end = finding["start"], finding["end"]
    return text[:start] + replacement_for(finding, treatment, pseudonym_map) + text[end:]


def apply_findings(src: Path, extraction: dict, findings: list[dict], workdir: Path, dst: Path,
                   treatment: str = "redact", pseudonym_map: dict[str, str] | None = None,
                   codes: list[str] | None = None) -> dict:
    """Aplica los hallazgos con status accepted|edited. Devuelve stats."""
    if treatment not in TREATMENTS:
        raise ValueError(f"Tratamiento inválido: {treatment}")
    src, workdir, dst = Path(src), Path(workdir), Path(dst)
    todo = [f for f in findings if f.get("status") in ("accepted", "edited")]
    has_email_finding = any(f["entity_code"] == "EMAIL" for f in todo)
    if treatment == "pseudonymize" and pseudonym_map is None:
        pseudonym_map = build_pseudonym_map(todo)
    stats: dict = {}
    for f in todo:
        stats[f["entity_code"]] = stats.get(f["entity_code"], 0) + 1

    if extraction["kind"] == "paged":
        # Page PNGs are derived cache files. Recreate a missing or incomplete
        # image from the retained original before applying a treatment.
        for page in extraction['pages']:
            if not page_image_ok(workdir / page['image']):
                if render_page(src,workdir,page['n']) is None:
                    raise EngineError('CORRUPT_FILE','No se pudo reconstruir la página para aplicar el tratamiento')
        _apply_paged(extraction, todo, workdir, dst, treatment, pseudonym_map)
        stats["pages"] = len(extraction["pages"])
    else:
        by_loc: dict[str, list] = {}
        for f in todo:
            if f.get("location") is not None:
                by_loc.setdefault(f["location"], []).append(f)
        fmt = extraction["format"]
        if fmt in ("doc", "xls"):
            conv = Path(extraction["converted"])
            clean = workdir / ("clean" + conv.suffix)
            if fmt == "doc":
                _apply_docx(conv, by_loc, clean, treatment, pseudonym_map, codes, has_email_finding)
            else:
                _apply_xlsx(conv, by_loc, clean, treatment, pseudonym_map)
            back = _convert(clean, fmt, workdir)
            shutil.copy(back, dst)
        elif fmt == "docx":
            _apply_docx(src, by_loc, dst, treatment, pseudonym_map, codes, has_email_finding)
        elif fmt == "xlsx":
            _apply_xlsx(src, by_loc, dst, treatment, pseudonym_map)
        else:
            text = extraction["texts"]["txt"]
            for f in sorted(by_loc.get("txt", []), key=lambda f: -f["start"]):
                text = _replace_span(text, f, treatment, pseudonym_map)
            dst.write_text(text, encoding="utf-8")
    stats["treatment"] = treatment
    stats["engine_version"] = engine.ENGINE_VERSION
    stats["total"] = sum(v for k, v in stats.items() if k in engine.ENTITY_TYPES)
    return stats


def protect(src: Path, dst: Path, workdir: Path, codes: list[str], treatment: str = "redact",
            analysis_level: str = "fast", pseudonyms=build_pseudonym_map) -> tuple[dict, dict]:
    """Ruta sin revisión humana (API directa y trabajos masivos): extract → detect → apply → verify.
    `pseudonyms(findings)` entrega el mapa de alias. Devuelve (stats, verificación)."""
    extraction = extract(src, workdir, analysis_level)
    findings = detect_findings(extraction, codes, workdir)
    pseudonym_map = pseudonyms(findings) if treatment == "pseudonymize" else None
    stats = apply_findings(src, extraction, findings, workdir, dst, treatment, pseudonym_map, codes)
    return stats, verify(dst, [f["text"] for f in findings if f.get("text")])


def _apply_paged(extraction, todo, workdir: Path, dst: Path, treatment: str,
                 pseudonym_map: dict[str, str] | None):
    import img2pdf
    from functools import lru_cache
    from statistics import median
    from PIL import Image, ImageDraw, ImageFont
    by_page: dict[int, list] = {}
    for f in todo:
        by_page.setdefault(f["page"], []).append(f)

    @lru_cache(maxsize=128)
    def replacement_font(size: int):
        candidates = (
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/System/Library/Fonts/Supplemental/Arial.ttf",
            "/Library/Fonts/Arial.ttf",
        )
        for candidate in candidates:
            if Path(candidate).exists():
                return ImageFont.truetype(candidate, size)
        try:
            return ImageFont.load_default(size=size)
        except TypeError:  # Pillow 10.0
            return ImageFont.load_default()

    def draw_page(p) -> "Image":
        img = Image.open(workdir / p["image"]).convert("RGB")
        draw = ImageDraw.Draw(img)
        replacements = []
        word_heights = [w[4] - w[2] for w in p.get("words", []) if w[4] > w[2]]
        normal_height = median(word_heights) if word_heights else None
        for f in by_page.get(p["n"], []):
            boxes = f["boxes"] or []
            if not boxes:
                continue
            if treatment == "redact":
                for x0, y0, x1, y1 in boxes:
                    h = y1 - y0
                    pad = max(1, int(h * 0.12)) if h < 80 else 2
                    draw.rectangle([x0 - pad, y0 - pad, x1 + pad, y1 + pad], fill="black")
                continue
            # Erase each word independently. A union spanning wrapped lines also
            # erases unrelated prose between the last word and the next line.
            for x0, y0, x1, y1 in boxes:
                pad = max(1, min(3, int((y1 - y0) * 0.12)))
                draw.rectangle([x0 - pad, y0 - pad, x1 + pad, y1 + pad], fill="white")
            # Place one substitute in a contiguous segment of an actual line.
            # The longest segment accommodates labels without inflating the font
            # to the height of the entire multiline finding.
            segments = _replacement_segments(boxes)
            x0, y0, x1, y1 = max(segments, key=lambda b: b[2] - b[0])
            value = replacement_for(f, treatment, pseudonym_map)
            height = min(y1 - y0, normal_height) if normal_height else y1 - y0
            # Invisible text layers can give words boxes much taller than their
            # visible ink. Keep substitutes at a readable body-text scale even
            # when most boxes on a scanned page share that inflated height.
            size = max(1, int(min(height * 0.8, p['width'] / 50)))
            while True:
                font = replacement_font(size)
                bounds = draw.textbbox((0, 0), value, font=font)
                if size <= 1 or (bounds[2] - bounds[0] <= max(x1 - x0, 1)
                                 and bounds[3] - bounds[1] <= max(y1 - y0, 1)):
                    break
                size -= 1
            replacements.append((value, font, (x0, y0, x1, y1), bounds))
        # Erase all originals before drawing any labels, including overlapping
        # findings from alternate OCR readings.
        for value, font, (x0, y0, x1, y1), bounds in replacements:
            text_height = bounds[3] - bounds[1]
            draw.text((x0 - bounds[0], y0 + (y1 - y0 - text_height) / 2 - bounds[1]),
                      value, fill="black", font=font)
        return img

    fmt = extraction["format"]
    pages = extraction["pages"]
    if fmt == "pdf":
        # Streaming: cada página se dibuja, se escribe como JPEG a disco y se libera.
        tmpdir = Path(tempfile.mkdtemp(prefix="apply_", dir=workdir))
        try:
            paths = []
            for p in pages:
                img = draw_page(p)
                jp = tmpdir / f"j{p['n']:05d}.jpg"
                img.save(jp, format="JPEG", quality=85)
                img.close()
                paths.append(str(jp))
            with dst.open("wb") as fh:
                img2pdf.convert(paths, outputstream=fh, pdfa=engine.SRGB_ICC)  # PDF/A-1b
        finally:
            shutil.rmtree(tmpdir, ignore_errors=True)
    elif fmt == "tiff" and len(pages) > 1:
        # El TIFF multipágina requiere todos los frames en memoria para save_all;
        # queda acotado por MAX_PAGES.
        frames = [draw_page(p) for p in pages]
        frames[0].save(dst, save_all=True, append_images=frames[1:], compression="tiff_deflate")
    else:
        img = draw_page(pages[0])
        img.save(dst, quality=90) if fmt == "jpg" else img.save(dst)


def _replacement_segments(boxes: list) -> list[tuple]:
    """Group neighboring boxes on the same visual line without bridging columns."""
    segments = []
    for box in sorted(boxes, key=lambda b: (b[1], b[0])):
        x0, y0, x1, y1 = box
        if x1 <= x0 or y1 <= y0:
            continue
        for index, segment in enumerate(segments):
            a0, b0, a1, b1 = segment
            height = min(y1 - y0, b1 - b0)
            overlap = min(y1, b1) - max(y0, b0)
            gap = max(x0 - a1, a0 - x1, 0)
            if overlap >= height * 0.5 and gap <= height * 1.5:
                segments[index] = (min(a0, x0), min(b0, y0), max(a1, x1), max(b1, y1))
                break
        else:
            segments.append(tuple(box))
    return segments


def _replace_text_nodes(ts, findings, treatment, pseudonym_map):
    """Reemplaza rangos desde el final conservando los runs no afectados."""
    for finding in sorted(findings, key=lambda item: -item["start"]):
        starts, pos = [], 0
        for node in ts:
            starts.append(pos)
            pos += len(node.text or "")
        start, end = finding["start"], finding["end"]
        affected = [i for i, node_start in enumerate(starts)
                    if node_start < end and node_start + len(ts[i].text or "") > start]
        if not affected:
            continue
        first, last = affected[0], affected[-1]
        first_text, last_text = ts[first].text or "", ts[last].text or ""
        prefix = first_text[:max(0, start - starts[first])]
        suffix = last_text[max(0, end - starts[last]):]
        ts[first].text = prefix + replacement_for(finding, treatment, pseudonym_map) + (suffix if first == last else "")
        for i in range(first + 1, last):
            ts[i].text = ""
        if last != first:
            ts[last].text = suffix


def _apply_docx(src: Path, by_loc: dict, dst: Path, treatment: str,
                pseudonym_map: dict[str, str] | None, codes: list[str] | None = None,
                has_email_finding: bool = False):
    import docx
    d = docx.Document(str(src))
    for loc, ts in engine._docx_paras(d):
        fs = by_loc.get(loc)
        if not fs:
            continue
        _replace_text_nodes(ts, fs, treatment, pseudonym_map)
    # El destino `mailto:` del hipervínculo no es texto del documento: hay que borrarlo
    # aparte o el correo sigue ahí aunque se vea tachado.
    # También cubre un correo agregado manualmente, aunque la selección original no
    # hubiera incluido EMAIL. Si el hallazgo se protegerá, su destino mailto no puede
    # quedar expuesto dentro del paquete DOCX.
    if codes is None or "EMAIL" in set(codes) or has_email_finding:
        engine._docx_strip_mailto(d)
    engine._docx_clean_props(d)
    d.save(str(dst))


def _apply_xlsx(src: Path, by_loc: dict, dst: Path, treatment: str,
                pseudonym_map: dict[str, str] | None):
    import openpyxl
    wb = openpyxl.load_workbook(str(src))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.comment is not None:
                    cell.comment = None
                if not isinstance(cell.value, str):
                    continue
                fs = by_loc.get(f"{ws.title}!{cell.coordinate}")
                if not fs:
                    continue
                v = cell.value
                for f in sorted(fs, key=lambda item: -item["start"]):
                    v = _replace_span(v, f, treatment, pseudonym_map)
                cell.value = v
    for attr in ("creator", "lastModifiedBy", "title", "subject", "description", "keywords"):
        setattr(wb.properties, attr, "")  # con None, openpyxl escribe "openpyxl" como creator
    wb.save(str(dst))


# ---------------------------------------------------------------------------
# verify  (FR-56 reducido: re-extracción + búsqueda normalizada + metadatos)
# ---------------------------------------------------------------------------

def _norm(s: str) -> str:
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9@]", "", s.lower())


def _norm_soft(s: str) -> str:
    """Minúsculas, sin tildes, separadores colapsados a un espacio."""
    s = "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9@]+", " ", s.lower()).strip()


def verify(dst: Path, suppressed: list[str]) -> dict:
    dst = Path(dst)
    checks = []
    ext = dst.suffix.lower()
    with tempfile.TemporaryDirectory() as td:
        e2 = extract(dst, Path(td))
        if e2["kind"] == "paged":
            # las páginas de un PDF protegido son imagen → siempre OCR del resultado.
            # Piezas = líneas de texto de cada página.
            pieces = [line for p in e2["pages"]
                      for line in __import__("backend.engine", fromlist=["x"])._words_text_map(p["words"])[0].split("\n")]
        else:
            pieces = list(e2["texts"].values())
    # Dos formas normalizadas, siempre por pieza (celda/párrafo/línea, unidas con \n
    # para que nada "cruce" piezas):
    #  A) suave: separadores colapsados a un espacio — literal aunque cambie el espaciado.
    #  B) pelada: sin separadores — atrapa reformateos (12.345.678-5 → 12345678-5), pero
    #     solo para valores largos (≥9): con valores cortos produce colisiones
    #     ("192.0.2.6" pelado es "192026", igual que la fecha "19/2026").
    corpus_soft = "\n".join(_norm_soft(t) for t in pieces)
    corpus_hard = "\n".join(_norm(t) for t in pieces)
    leaked = sorted({v for v in suppressed
                     if (len(_norm(v)) >= 4 and _norm_soft(v) and _norm_soft(v) in corpus_soft)
                     or (len(_norm(v)) >= 9 and _norm(v) in corpus_hard)})
    checks.append({"name": "sin_valores_suprimidos", "ok": not leaked,
                   "detail": f"{len(suppressed)} valores buscados" + (f"; FILTRADOS: {leaked[:5]}" if leaked else "")})

    meta_ok, meta_detail = True, "sin metadatos de autor"
    try:
        if ext == ".pdf":
            import pdfplumber
            with pdfplumber.open(dst) as pdf:
                m = pdf.metadata or {}
            bad = {k: v for k, v in m.items() if k in ("Author", "Title", "Subject", "Keywords") and v}
            meta_ok, meta_detail = not bad, str(bad) if bad else meta_detail
        elif ext == ".docx":
            import docx
            cp = docx.Document(str(dst)).core_properties
            bad = {a: getattr(cp, a) for a in ("author", "last_modified_by", "title") if getattr(cp, a)}
            meta_ok, meta_detail = not bad, str(bad) if bad else meta_detail
        elif ext == ".xlsx":
            import openpyxl
            pr = openpyxl.load_workbook(str(dst)).properties
            bad = {a: getattr(pr, a) for a in ("creator", "lastModifiedBy", "title") if getattr(pr, a)}
            meta_ok, meta_detail = not bad, str(bad) if bad else meta_detail
    except Exception as exc:  # noqa: BLE001
        meta_ok, meta_detail = False, f"no se pudo verificar: {exc}"
    checks.append({"name": "metadatos_limpios", "ok": meta_ok, "detail": meta_detail})

    return {"ok": all(c["ok"] for c in checks), "checks": checks}
